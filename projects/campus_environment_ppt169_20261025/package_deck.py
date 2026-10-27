# -*- coding: utf-8 -*-
"""Create the embedded and linked campus deck deliverables."""

from __future__ import annotations

import json
import os
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation


ROOT = Path(__file__).resolve().parent
BUILD = ROOT / "_deck_build.pptx"
EMBEDDED_SRC = ROOT / "_deck_animated.pptx"
OUT = ROOT / "deliverables"
EMBEDDED_NAME = "校园环境_三个角落_同一个问题_10页_嵌入版.pptx"
LINK_DIR = ROOT / "_linked_stage"
LINK_NAME = "校园环境_三个角落_同一个问题_10页_链接版.pptx"
MEDIA_DIR = LINK_DIR / "media"
FONT_DIR = LINK_DIR / "fonts"
PREVIEW_SRC = ROOT / "preview_outlined"

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def safe_rmtree(path: Path) -> None:
    target = path.resolve()
    root = ROOT.resolve()
    if target == root or root not in target.parents:
        raise RuntimeError(f"Refusing to delete outside project: {target}")
    shutil.rmtree(target)


def load_slots() -> dict[int, str]:
    records = json.loads((ROOT / "video_slots.json").read_text(encoding="utf-8"))
    return {int(item["slide"]): Path(item["video"]).name for item in records}


def make_linked(src: Path, dst: Path, slots: dict[int, str]) -> None:
    with zipfile.ZipFile(src) as zin:
        names = zin.namelist()
        replacements: dict[str, bytes] = {}
        external_count = 0
        for slide_no, video_name in slots.items():
            slide_part = f"ppt/slides/slide{slide_no}.xml"
            rels_part = f"ppt/slides/_rels/slide{slide_no}.xml.rels"
            if slide_part not in names or rels_part not in names:
                raise RuntimeError(f"Missing slide part for video on slide {slide_no}")
            target_uri = (MEDIA_DIR / video_name).resolve().as_uri()

            rels = etree.fromstring(zin.read(rels_part))
            linked_ids = []
            for rel in rels.findall(f"{{{REL_NS}}}Relationship"):
                rel_type = rel.get("Type", "")
                if rel_type.endswith("/video") or rel_type.endswith("/media"):
                    rel.set("Target", target_uri)
                    rel.set("TargetMode", "External")
                    linked_ids.append(rel.get("Id"))
            if not linked_ids:
                raise RuntimeError(f"No media relationship on slide {slide_no}")
            external_count += len(linked_ids)

            slide = etree.fromstring(zin.read(slide_part))
            for element in slide.iter():
                if element.tag.endswith("}videoFile"):
                    rid = element.get(f"{{{R_NS}}}link")
                    if rid not in linked_ids:
                        raise RuntimeError(f"Video link on slide {slide_no} is not external")
                if element.tag.endswith("}media"):
                    embed_id = element.get(f"{{{R_NS}}}embed")
                    if embed_id:
                        del element.attrib[f"{{{R_NS}}}embed"]
                    element.set(f"{{{R_NS}}}link", linked_ids[0])
            replacements[slide_part] = etree.tostring(
                slide, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            replacements[rels_part] = etree.tostring(
                rels, xml_declaration=True, encoding="UTF-8", standalone=True
            )

        tmp = dst.with_name(dst.stem + ".tmp" + dst.suffix)
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                if info.filename.startswith("ppt/media/") and info.filename.lower().endswith(".mp4"):
                    continue
                data = replacements.get(info.filename, zin.read(info.filename))
                zout.writestr(info, data)
        os.replace(tmp, dst)
    print(f"linked: {external_count} external media relationships")


def copy_media_and_fonts() -> None:
    MEDIA_DIR.mkdir(parents=True, exist_ok=True)
    FONT_DIR.mkdir(parents=True, exist_ok=True)
    for name in ("video1.mp4", "video2.mp4", "video3.mp4"):
        shutil.copy2(ROOT / "images" / name, MEDIA_DIR / name)
    fonts = {
        Path(r"C:\Windows\Fonts\simkai.ttf"): FONT_DIR / "KaiTi_simkai.ttf",
        Path(r"C:\Users\camel\AppData\Local\Microsoft\Windows\Fonts\STZHONGS.TTF"): FONT_DIR / "STZhongsong.ttf",
        Path(r"C:\Users\camel\AppData\Local\Microsoft\Windows\Fonts\TrajanProDeck.ttf"): FONT_DIR / "TrajanProDeck.ttf",
    }
    for src, dst in fonts.items():
        if not src.exists():
            raise FileNotFoundError(src)
        shutil.copy2(src, dst)


def write_readme() -> None:
    text = """校园环境汇报 / 两种交付版本

1. 嵌入版
   单个 PPTX，3 段视频与 Trajan Pro Deck、楷体、华文中宋均已嵌入。
   适合直接发送、拷走或现场投影。

2. 链接版（本压缩包内）
   PPTX 中的 3 段视频链接到同目录下的 media 文件夹；字体仍嵌入在 PPTX 内。
   解压后保持以下同级结构，不能只取出 PPTX：

   交付说明.txt
   校园环境_三个角落_同一个问题_10页_链接版.pptx
   media/
     video1.mp4
     video2.mp4
     video3.mp4
   fonts/
     KaiTi_simkai.ttf
     STZhongsong.ttf
     TrajanProDeck.ttf
   修复链接.cmd
   修复链接.ps1

   解压到任意位置后，如果视频不播放，双击“修复链接.cmd”，会把视频链接重新
   指向当前文件夹内的 media。也可以直接在 PowerPoint 中右键视频，选择“更改媒体”。

   链接版的视频封面和版式与嵌入版一致；视频本身不进入 PPTX，因此文件更小。

字体说明
  大标题已转为高分辨率轮廓图（Trajan Pro），实际放映不依赖字体安装或加载。
  引语与副标题为楷体，正文、英文和数字为华文中宋。
  两版 PPTX 同时保留字体嵌入；fonts 文件夹中的 TTF 仅作为备用安装文件。

动效说明
  全片自动淡入，每项 0.4 秒、60 毫秒错峰，无需逐项点击。
  页面之间使用柔淡切换；背景、渐变底与大幅圆弧保持静止，避免整页闪白。
  照片上的渐淡罩层与对应照片同步出现。

内容说明
   10 页，16:9，英文汇报；不含演讲提示、时间戳和逐字稿。
   14 张照片与 3 段视频全部使用。
"""
    (LINK_DIR / "交付说明.txt").write_text(text, encoding="utf-8-sig")

    ps1 = r"""$ErrorActionPreference = 'Stop'
$root = Split-Path -Parent $MyInvocation.MyCommand.Path
$pptx = Get-ChildItem -LiteralPath $root -Filter '*.pptx' | Select-Object -First 1
if (-not $pptx) { throw 'No PPTX found next to this script.' }
$media = Join-Path $root 'media'
if (-not (Test-Path -LiteralPath $media)) { throw 'media folder not found.' }

Add-Type -AssemblyName System.IO.Compression
Add-Type -AssemblyName System.IO.Compression.FileSystem

$zip = [System.IO.Compression.ZipFile]::Open($pptx.FullName, 'Update')
try {
    $names = @($zip.Entries |
        Where-Object { $_.FullName -match '^ppt/slides/_rels/slide\d+\.xml\.rels$' } |
        ForEach-Object { $_.FullName })
    foreach ($name in $names) {
        $entry = $zip.GetEntry($name)
        $reader = New-Object System.IO.StreamReader($entry.Open(), [Text.Encoding]::UTF8)
        $text = $reader.ReadToEnd()
        $reader.Dispose()

        $doc = New-Object System.Xml.XmlDocument
        $doc.LoadXml($text)
        $ns = New-Object System.Xml.XmlNamespaceManager($doc.NameTable)
        $ns.AddNamespace('r', 'http://schemas.openxmlformats.org/package/2006/relationships')
        $changed = $false
        foreach ($rel in $doc.SelectNodes('//r:Relationship', $ns)) {
            $type = $rel.GetAttribute('Type')
            if ($type.EndsWith('/video') -or $type.EndsWith('/media')) {
                $leaf = Split-Path -Leaf $rel.GetAttribute('Target')
                $path = Join-Path $media $leaf
                if (Test-Path -LiteralPath $path) {
                    $rel.SetAttribute('Target', ([Uri]$path).AbsoluteUri)
                    $rel.SetAttribute('TargetMode', 'External')
                    $changed = $true
                }
            }
        }
        if ($changed) {
            $entry.Delete()
            $newEntry = $zip.CreateEntry($name, [System.IO.Compression.CompressionLevel]::Optimal)
            $stream = $newEntry.Open()
            $writer = New-Object System.IO.StreamWriter($stream, (New-Object System.Text.UTF8Encoding($false)))
            $writer.Write($doc.OuterXml)
            $writer.Dispose()
            $stream.Dispose()
        }
    }
}
finally {
    $zip.Dispose()
}
Write-Host 'Video links updated to:' $media
"""
    (LINK_DIR / "修复链接.ps1").write_text(ps1, encoding="utf-8-sig")
    (LINK_DIR / "修复链接.cmd").write_text(
        '@echo off\r\nchcp 65001 >nul\r\npowershell -NoProfile -ExecutionPolicy Bypass -File "%~dp0修复链接.ps1"\r\npause\r\n',
        encoding="utf-8",
    )


def copy_previews() -> None:
    out = OUT / "preview"
    out.mkdir(parents=True, exist_ok=True)
    for i in range(1, 11):
        shutil.copy2(PREVIEW_SRC / f"slide-{i}.png", out / f"slide-{i}.png")
    images = [Image.open(out / f"slide-{i}.png").convert("RGB") for i in range(1, 11)]
    width, height = images[0].size
    scale = 0.42
    tw, th = int(width * scale), int(height * scale)
    contact = Image.new("RGB", (tw * 2 + 36, th * 5 + 54), (236, 236, 232))
    for i, image in enumerate(images):
        contact.paste(
            image.resize((tw, th), Image.Resampling.LANCZOS),
            ((i % 2) * (tw + 12) + 12, (i // 2) * (th + 10) + 10),
        )
    contact.save(out / "联系表.jpg", quality=90)


def audit(path: Path, video_parts: int) -> None:
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        slides = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        videos = [n for n in names if n.lower().endswith(".mp4")]
        fonts = [n for n in names if n.startswith("ppt/fonts/")]
        if len(slides) != 10:
            raise RuntimeError(f"{path.name}: expected 10 slides, found {len(slides)}")
        if len(videos) != video_parts:
            raise RuntimeError(f"{path.name}: expected {video_parts} videos, found {len(videos)}")
        if len(fonts) != 3:
            raise RuntimeError(f"{path.name}: expected 3 embedded fonts, found {len(fonts)}")
        timing_count = 0
        click_effects = 0
        for slide_name in slides:
            root = etree.fromstring(z.read(slide_name))
            timing = root.find(f"{{{P_NS}}}timing")
            transition = root.find(f"{{{P_NS}}}transition")
            if timing is None or transition is None:
                raise RuntimeError(f"{path.name}: missing timing/transition in {slide_name}")
            timing_count += 1
            click_effects += sum(
                1
                for node in timing.iter(f"{{{P_NS}}}cTn")
                if node.get("presetID") == "10" and node.get("nodeType") == "clickEffect"
            )
        if timing_count != 10 or click_effects:
            raise RuntimeError(
                f"{path.name}: expected 10 automatic timings, found {timing_count}, clicks {click_effects}"
            )
        if video_parts == 0:
            external = 0
            for name in names:
                if name.startswith("ppt/slides/_rels/") and name.endswith(".rels"):
                    root = etree.fromstring(z.read(name))
                    external += sum(
                        1
                        for rel in root.findall(f"{{{REL_NS}}}Relationship")
                        if rel.get("TargetMode") == "External"
                        and rel.get("Type", "").rsplit("/", 1)[-1] in {"video", "media"}
                    )
            if external != 6:
                raise RuntimeError(f"{path.name}: expected 6 external media links, found {external}")
    prs = Presentation(path)
    if len(prs.slides) != 10:
        raise RuntimeError(f"{path.name}: python-pptx sees {len(prs.slides)} slides")
    movie_names = {
        shape.name
        for slide in prs.slides
        for shape in slide.shapes
        if shape.name.startswith("FIELD_VIDEO_")
    }
    if movie_names != {"FIELD_VIDEO_1", "FIELD_VIDEO_2", "FIELD_VIDEO_3"}:
        raise RuntimeError(f"{path.name}: movie shapes mismatch: {movie_names}")
    print(f"audit ok: {path.name} ({path.stat().st_size:,} bytes)")


def main() -> None:
    if not BUILD.exists() or not EMBEDDED_SRC.exists():
        raise FileNotFoundError("Run build_deck.py, embed_fonts.py and animate_deck.py first.")
    if not PREVIEW_SRC.exists():
        raise FileNotFoundError(f"Preview folder missing: {PREVIEW_SRC}")
    OUT.mkdir(parents=True, exist_ok=True)
    if LINK_DIR.exists():
        safe_rmtree(LINK_DIR)
    LINK_DIR.mkdir(parents=True, exist_ok=True)
    embedded = OUT / EMBEDDED_NAME
    shutil.copy2(EMBEDDED_SRC, embedded)

    copy_media_and_fonts()
    linked = LINK_DIR / LINK_NAME
    make_linked(EMBEDDED_SRC, linked, load_slots())
    write_readme()
    copy_previews()

    audit(embedded, video_parts=3)
    audit(linked, video_parts=0)

    archive_base = OUT / "校园环境_三个角落_同一个问题_10页_链接版"
    archive_path = archive_base.with_suffix(".zip")
    if archive_path.exists():
        archive_path.unlink()
    archive = shutil.make_archive(str(archive_base), "zip", root_dir=LINK_DIR)
    print("archive:", archive)

    for stale in (
        OUT / "linked_package",
        OUT / "交付说明.txt",
        OUT / "校园环境_10页_链接版.zip",
        archive_base,
    ):
        if stale.is_dir():
            safe_rmtree(stale)
        elif stale.exists():
            stale.unlink()
    safe_rmtree(LINK_DIR)
    print("finalised:", OUT)


if __name__ == "__main__":
    main()
