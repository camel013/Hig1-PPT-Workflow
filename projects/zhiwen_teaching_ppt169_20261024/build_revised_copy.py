# -*- coding: utf-8 -*-
"""Rebuild the existing deck with the revised explanatory copy.

The visual builders stay untouched.  Text calls are intercepted at runtime and
replaced only by page, so quotation blocks, geometry, colour and type styling
remain exactly as they were in the approved deck.
"""
from __future__ import annotations

import pathlib
import re
import sys


HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import deck_lib as D
import state
from pptx.oxml.ns import qn


OLD_COPY = HERE / "蛭文_教学读本_95页_讲解文字稿.txt"
NEW_COPY = HERE / "蛭文_教学读本_95页_讲解文字稿_去AI改写稿.txt"
OUT = HERE / "蛭文_教学读本_95页_新讲解版.pptx"


def _pages(path: pathlib.Path) -> dict[int, list[str]]:
    text = path.read_text(encoding="utf-8")
    chunks = re.split(r"(?=^## 第\d{2}页$)", text, flags=re.M)
    pages: dict[int, list[str]] = {}
    for chunk in chunks:
        match = re.match(r"## 第(\d{2})页\s*", chunk)
        if not match:
            continue
        page = int(match.group(1))
        pages[page] = [
            para.strip()
            for para in chunk[match.end():].split("\n\n")
            if para.strip() and para.strip() != "---"
        ]
    return pages


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", text)


OLD = _pages(OLD_COPY)
NEW = _pages(NEW_COPY)

DIRECT: dict[int, dict[str, str]] = {}
SEQUENCES: dict[int, list[tuple[str, str]]] = {}
for page in range(1, 96):
    mapping: dict[str, str] = {}
    for old, new in zip(OLD[page], NEW[page]):
        mapping[_norm(old)] = new
    DIRECT[page] = mapping

    sequences: list[tuple[str, str]] = []
    for start in range(len(OLD[page])):
        old_concat = ""
        new_concat: list[str] = []
        for end in range(start, len(OLD[page])):
            old_concat += _norm(OLD[page][end])
            new_concat.append(NEW[page][end])
            if end > start:
                sequences.append((old_concat, "".join(new_concat)))
    SEQUENCES[page] = sequences


# Two visual text boxes on page 17 carry one sentence split over two lines.
SPECIAL = {
    5: {
        "序位于正文与读者之间，像一道门槛：它不讲故事，却规定我们如何进入故事。"
        "热奈特把标题、序、跋、封面、题辞这类文字称为“副文本”——正文的门厅。":
            "序站在正文与读者之间，像一道门槛。它不承担叙事，却预先规定读者如何进入叙事。"
            "热奈特把标题、序、跋、封面、题辞这类文字称为“副文本”，也就是正文的门厅。",
    },
    15: {
        "先把它放回它借来的那个“壳”里，再看它换了什么心。":
            "先回到它借来的“壳”，再看它换掉了什么。",
    },
    17: {
        "“有声”不是热闹，": "“有声”意味着无法退出，",
        "是另一种不能退出。": "热闹只是表面。",
    },
}


CURRENT_PAGE = 0
STATS = {"replaced": 0, "unmatched_changed": 0}
GROUP_COUNTER = 0
TEXT_DEPTH = 0


def lookup(text: str) -> str:
    key = _norm(text)
    page = CURRENT_PAGE
    if not page:
        return text

    special = SPECIAL.get(page, {})
    if key in special:
        return special[key]

    replacement = DIRECT.get(page, {}).get(key)
    if replacement is not None:
        if _norm(replacement) != key:
            STATS["replaced"] += 1
        return replacement

    # Some sections draw two paragraphs in one call. Join their replacements
    # without adding punctuation or rephrasing either paragraph.
    for old_concat, new_concat in SEQUENCES.get(page, []):
        if key == old_concat:
            return new_concat
    return text


def install() -> None:
    global GROUP_COUNTER
    original_new_slide = D.new_slide
    original_text = D.text
    original_vline = D.vline
    original_rich = D.rich

    def tag_shapes(slide, before: int, kind: str) -> None:
        global GROUP_COUNTER
        token = f"animgrp:{CURRENT_PAGE:02d}:{GROUP_COUNTER:03d}:{kind}"
        GROUP_COUNTER += 1
        for shape in list(slide.shapes)[before:]:
            c_nv = shape._element.find(".//" + qn("p:cNvPr"))
            if c_nv is not None:
                c_nv.set("descr", token)

    def new_slide(*args, **kwargs):
        global CURRENT_PAGE, GROUP_COUNTER
        CURRENT_PAGE = state.PAGE
        GROUP_COUNTER = 0
        return original_new_slide(*args, **kwargs)

    def text(slide, x, y, w, value, *args, **kwargs):
        global TEXT_DEPTH
        if TEXT_DEPTH:
            return original_text(slide, x, y, w, value, *args, **kwargs)
        font = kwargs.get("font", args[1] if len(args) > 1 else D.SONG)
        kind = "body" if font == D.SONG else "heading"
        before = len(slide.shapes)
        TEXT_DEPTH = 1
        try:
            result = original_text(slide, x, y, w, lookup(value), *args, **kwargs)
        finally:
            TEXT_DEPTH = 0
        tag_shapes(slide, before, kind)
        return result

    def vline(slide, *args, **kwargs):
        before = len(slide.shapes)
        result = original_vline(slide, *args, **kwargs)
        tag_shapes(slide, before, "heading")
        return result

    def rich(slide, *args, **kwargs):
        before = len(slide.shapes)
        result = original_rich(slide, *args, **kwargs)
        tag_shapes(slide, before, "body")
        return result

    D.new_slide = new_slide
    D.text = text
    D.vline = vline
    D.rich = rich


def main() -> None:
    install()
    import deck_front
    import deck_unit1
    import deck_unit2
    import deck_unit3
    import deck_unit4
    import deck_end

    deck_front.build()
    deck_unit1.unit1()
    deck_unit2.unit2()
    deck_unit3.unit3()
    deck_unit4.unit4()
    deck_end.ending()
    D.add_color_chords()
    D.prs.save(OUT)
    print("slides", len(D.prs.slides), "replaced", STATS["replaced"])
    print(OUT)


if __name__ == "__main__":
    main()
