# -*- coding: utf-8 -*-
from pptx import Presentation
from collections import OrderedDict
p=Presentation(r'projects\zhiwen_teaching_ppt169_20261024\蛭文_教学读本_95页_新讲解版.pptx')
for si,s in enumerate(p.slides,1):
    groups=OrderedDict()
    for sh in s.shapes:
        cn=sh._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        d=cn.get('descr','') if cn is not None else ''
        if not d.startswith('animgrp:'): continue
        groups.setdefault(d,[]).append(sh)
    print(f'=== {si:02d} groups={len(groups)} ===')
    for d,arr in groups.items():
        t=' / '.join((getattr(x,'text','').replace('\n','|')[:24] or '[shape]') for x in arr)
        print(d.split(':')[-2], d.split(':')[-1], len(arr), t[:100])
