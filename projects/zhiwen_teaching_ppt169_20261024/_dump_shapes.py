# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
PX=9600
p=Presentation(r'projects\zhiwen_teaching_ppt169_20261024\蛭文_教学读本_95页_新讲解版.pptx')
pages={int(x) for x in sys.argv[1:]}
for si,s in enumerate(p.slides,1):
    if si not in pages: continue
    print(f'\n=== {si:02d} ===')
    for sh in s.shapes:
        cn=sh._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        d=cn.get('descr','') if cn is not None else ''
        g=d.split(':')[-2] if d.startswith('animgrp:') else ''
        t=getattr(sh,'text','').replace('\n',' | ')
        print(f'{sh.shape_id:3} g{g:>3} {sh.name[:22]:22} [{sh.left/PX:6.0f},{sh.top/PX:6.0f},{sh.width/PX:6.0f},{sh.height/PX:6.0f}] {t[:64]}')
