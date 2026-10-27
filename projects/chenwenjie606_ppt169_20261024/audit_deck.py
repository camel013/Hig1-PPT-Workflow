from __future__ import annotations

from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu


ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "陈文杰_606分之后_9页.pptx"


def main() -> None:
    prs = Presentation(PPTX)
    slide_w = prs.slide_width
    slide_h = prs.slide_height
    fonts = Counter()
    sizes = Counter()
    bold_runs = 0
    total_runs = 0
    overflow = []
    styled = 0
    body_paragraphs = []
    empty_notes = []

    for slide_index, slide in enumerate(prs.slides, 1):
        if not slide.notes_slide.notes_text_frame.text.strip():
            empty_notes.append(slide_index)

        for shape in slide.shapes:
            if shape._element.find(qn("p:style")) is not None:
                styled += 1
            if shape.left + shape.width > slide_w + Emu(5000) or shape.top + shape.height > slide_h + Emu(5000):
                if shape.has_text_frame and shape.text.strip():
                    overflow.append(
                        (slide_index, shape.text.strip()[:24], "text-shape")
                    )
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                p_sizes = []
                for run in paragraph.runs:
                    if not run.text:
                        continue
                    total_runs += 1
                    if run.font.bold:
                        bold_runs += 1
                    fonts[run.font.name or "(none)"] += 1
                    size = round(run.font.size.pt, 1) if run.font.size else 0
                    sizes[size] += 1
                    p_sizes.append(size)
                if (
                    p_sizes
                    and 22 <= max(p_sizes) <= 26
                    and len(paragraph.text.strip()) >= 12
                ):
                    ppr = paragraph._p.pPr
                    indent = int(ppr.get("indent", "0")) if ppr is not None else 0
                    body_paragraphs.append(
                        (
                            slide_index,
                            paragraph.text.strip()[:26],
                            max(p_sizes),
                            indent,
                            paragraph.line_spacing,
                        )
                    )

    zero_indent = [row for row in body_paragraphs if row[3] == 0]
    wrong_spacing = [
        row for row in body_paragraphs if row[4] not in (1.5, 36.0)
    ]

    print(f"slides={len(prs.slides)}")
    print(f"size_bytes={PPTX.stat().st_size}")
    print(f"fonts={dict(fonts)}")
    print(f"bold={bold_runs}/{total_runs}")
    print(f"font_sizes={dict(sorted(sizes.items()))}")
    print(f"text_overflow={overflow}")
    print(f"remaining_styles={styled}")
    print(f"empty_notes={empty_notes}")
    print(f"body_paragraphs={len(body_paragraphs)}")
    print(f"body_zero_indent={zero_indent}")
    print(f"body_wrong_line_spacing={wrong_spacing}")


if __name__ == "__main__":
    main()
