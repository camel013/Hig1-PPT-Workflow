from __future__ import annotations

import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "陈文杰_606分之后_9页.pptx"

FONT = "Microsoft YaHei"
INK = RGBColor(0x11, 0x11, 0x11)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x0F, 0x4B, 0xD8)
PALE_BLUE = RGBColor(0xB8, 0xC8, 0xE8)

SLIDE_W = 13.333
SLIDE_H = 7.5
LEFT = 0.82
RIGHT = 12.52


def rgb_hex(color: RGBColor) -> str:
    return str(color)


def _set_run_font(run, size: float, color: RGBColor) -> None:
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)
    rpr.set("lang", "zh-CN")
    rpr.set("b", "1")


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str | list,
    size: float = 24,
    color: RGBColor = INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    line_spacing: float = 1.5,
    first_indent: bool = False,
    margins: float = 0,
    paragraphs: bool = True,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(margins)
    tf.margin_right = Inches(margins)
    tf.margin_top = Inches(margins)
    tf.margin_bottom = Inches(margins)
    tf.vertical_anchor = valign

    if isinstance(text, str):
        text = text.split("\n")
    for index, item in enumerate(text):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        if first_indent and align == PP_ALIGN.LEFT:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", "0")
            ppr.set("indent", str(int(size * 2 * 12700)))
        runs = item if isinstance(item, list) else [(item, color, size)]
        for run_text, run_color, run_size in runs:
            run = p.add_run()
            run.text = run_text
            _set_run_font(run, run_size, run_color)
    return box


def add_indented_lines(
    slide, x, y, w, h, lines: list[str], size: float = 24, color: RGBColor = INK
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.5
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    ppr = p._p.get_or_add_pPr()
    ppr.set("marL", "0")
    ppr.set("indent", str(int(size * 2 * 12700)))
    for i, line in enumerate(lines):
        run = p.add_run()
        run.text = line
        _set_run_font(run, size, color)
        if i < len(lines) - 1:
            p.add_line_break()
    return box


def add_rect(slide, x, y, w, h, color, line=None, line_width=0.75):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    return shape


def add_rule(slide, x1, y, x2, color=INK, width=0.9):
    return add_rect(slide, x1, y, x2 - x1, width / 72, color)


def add_vrule(slide, x, y1, y2, color=INK, width=0.9):
    return add_rect(slide, x, y1, width / 72, y2 - y1, color)


def add_circle(
    slide, x, y, diameter, color=BLUE, width=2.0, fill_color=None
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x),
        Inches(y),
        Inches(diameter),
        Inches(diameter),
    )
    if fill_color is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def add_dot(slide, x, y, diameter, color=BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x),
        Inches(y),
        Inches(diameter),
        Inches(diameter),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def cubic(p0, p1, p2, p3, count=90):
    points = []
    for i in range(count):
        t = i / (count - 1)
        u = 1 - t
        x = u**3 * p0[0] + 3 * u**2 * t * p1[0] + 3 * u * t**2 * p2[0] + t**3 * p3[0]
        y = u**3 * p0[1] + 3 * u**2 * t * p1[1] + 3 * u * t**2 * p2[1] + t**3 * p3[1]
        points.append((x, y))
    return points


def add_polyline(slide, points, color=BLUE, width=2.0):
    scale = 914400
    builder = slide.shapes.build_freeform(
        points[0][0] * scale, points[0][1] * scale, scale=1.0
    )
    builder.add_line_segments(
        [(x * scale, y * scale) for x, y in points[1:]], close=False
    )
    shape = builder.convert_to_shape()
    shape.fill.background()
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def add_kicker(slide, number: int, label: str, dark: bool = False):
    color = WHITE if dark else BLUE
    add_text(
        slide,
        LEFT,
        0.39,
        8.6,
        0.28,
        f"{number:02d} / 09    {label}",
        size=12,
        color=color,
        line_spacing=1.0,
    )
    add_rule(
        slide,
        LEFT,
        0.77,
        RIGHT,
        color=RGBColor(0x58, 0x65, 0x7C) if dark else PALE_BLUE,
        width=0.65,
    )


def add_page_number(slide, number: int, dark: bool = False):
    add_text(
        slide,
        12.05,
        7.11,
        0.5,
        0.24,
        f"{number:02d}",
        size=12,
        color=WHITE if dark else BLUE,
        align=PP_ALIGN.RIGHT,
        line_spacing=1.0,
    )


def add_source(slide, text: str, dark: bool = False, y: float = 7.1):
    add_text(
        slide,
        LEFT,
        y,
        9.6,
        0.24,
        text,
        size=12,
        color=RGBColor(0xD8, 0xE1, 0xF1) if dark else RGBColor(0x5E, 0x6B, 0x82),
        line_spacing=1.0,
    )


def add_title(slide, text: str, dark: bool = False, y: float = 0.99):
    add_text(
        slide,
        LEFT,
        y,
        10.8,
        0.68,
        text,
        size=34,
        color=WHITE if dark else INK,
        line_spacing=1.0,
    )


def add_notes(slide, narration: str):
    slide.notes_slide.notes_text_frame.text = narration


def white_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    return slide


def black_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = INK
    return slide


def strip_style_and_shadow(slide):
    for shape in slide.shapes:
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)
        try:
            shape.shadow.inherit = False
        except Exception:
            pass


def slide_01(prs):
    s = white_slide(prs)
    add_kicker(s, 1, "封面")
    add_circle(s, 9.45, 1.52, 5.55, color=BLUE, width=2.6)
    add_circle(s, 10.95, 3.10, 2.35, color=PALE_BLUE, width=0.9)
    add_text(s, 0.66, 1.56, 7.8, 2.5, "606", 168, BLUE, line_spacing=0.92)
    add_text(s, 5.73, 2.58, 3.5, 0.8, "分之后", 46, INK, line_spacing=1.0)
    add_rule(s, 0.82, 4.45, 8.15, color=INK, width=1.1)
    add_text(s, 0.84, 4.67, 7.1, 0.5, "陈文杰的录取时刻", 26, INK, line_spacing=1.0)
    add_text(s, 0.84, 5.28, 5.0, 0.4, "四川凉山 · 2026", 22, BLUE, line_spacing=1.0)
    add_text(s, 9.96, 5.86, 2.45, 0.7, "从凉山\n走向大学", 20, INK, line_spacing=1.25)
    add_source(s, "资料：四川观察 7 月 20 日发布画面")
    add_page_number(s, 1)
    add_notes(
        s,
        "今天讲陈文杰的故事。他来自四川凉山彝族自治州宁南县。2026 年 7 月 19 日，他等来一张四川大学的录取结果。",
    )
    return s


def slide_02(prs):
    s = white_slide(prs)
    add_kicker(s, 2, "人物")
    add_title(s, "一个凉山少年，等来结果")
    add_text(s, 0.82, 2.18, 3.38, 2.35, "606", 126, BLUE, line_spacing=0.92)
    add_text(s, 3.64, 3.74, 0.9, 0.45, "分", 26, INK, line_spacing=1.0)
    add_text(s, 0.86, 4.68, 3.0, 0.42, "高考成绩", 22, INK, line_spacing=1.0)
    add_vrule(s, 4.63, 2.15, 6.42, color=INK, width=1.0)

    facts = [
        ("2026.07.19", "录取日期", "日期"),
        ("本科批次高校专项计划", "报考通道", "项目"),
        ("四川大学 · 工科试验班（物理类）", "录取学校与专业", "结果"),
    ]
    ys = [2.18, 3.43, 4.68]
    for (main, note, tag), y in zip(facts, ys):
        add_text(s, 5.02, y, 0.48, 0.3, tag, 12, BLUE, line_spacing=1.0)
        add_text(s, 5.58, y - 0.02, 6.72, 0.58, main, 28, INK, line_spacing=1.0)
        add_text(s, 5.6, y + 0.55, 6.2, 0.34, note, 20, BLUE, line_spacing=1.0)
        if y < 4.6:
            add_rule(s, 5.02, y + 1.02, 12.3, PALE_BLUE, width=0.65)

    add_indented_lines(
        s,
        0.84,
        5.57,
        3.56,
        1.0,
        ["来自四川凉山", "彝族自治州宁南县。"],
        24,
        INK,
    )
    add_source(s, "画面记录：2026.07.19 · 发布：2026.07.20")
    add_page_number(s, 2)
    add_notes(
        s,
        "这张录取结果上，最醒目的是 606 分。它记录了陈文杰的高考成绩，也把他送进本科批次高校专项计划的录取流程。结果落在四川大学工科试验班物理类。",
    )
    return s


def slide_03(prs):
    s = white_slide(prs)
    add_kicker(s, 3, "成绩")
    add_title(s, "606分，他交出的答卷")
    add_text(s, 0.68, 1.48, 8.5, 4.2, "606", 230, INK, line_spacing=0.9)
    add_text(s, 8.68, 1.77, 3.22, 0.95, "他自己\n考出的成绩", 28, BLUE, line_spacing=1.25)
    add_circle(s, 9.11, 3.15, 4.4, color=BLUE, width=1.6)
    add_rule(s, 0.82, 5.32, 8.25, color=BLUE, width=1.8)
    add_indented_lines(
        s,
        0.86,
        5.58,
        7.45,
        1.1,
        ["606分不是凭空落下的数字。", "它来自陈文杰走过的备考路。"],
        24,
        INK,
    )
    add_text(s, 10.18, 5.85, 2.2, 0.42, "亲手交出", 18, BLUE, line_spacing=1.0)
    add_source(s, "资料：四川观察 7 月 20 日发布画面")
    add_page_number(s, 3)
    add_notes(
        s,
        "606 分不是凭空落下的数字，它来自陈文杰走过的备考路。公开画面没有展开那些日子，但分数本身，已经记录了他交出的答卷。",
    )
    return s


def slide_04(prs):
    s = white_slide(prs)
    add_kicker(s, 4, "选择")
    add_title(s, "这条路，他自己走完")
    add_text(
        s,
        0.84,
        1.67,
        8.1,
        0.98,
        "专项计划提供了机会，资格审核、志愿填报和录取，仍要由他一步步完成。",
        22,
        INK,
        line_spacing=1.5,
        first_indent=True,
    )
    pts = cubic((1.02, 5.92), (5.00, 5.72), (6.12, 2.82), (11.84, 2.38), 110)
    add_polyline(s, pts, BLUE, 2.4)
    path_nodes = [
        (pts[15], "资格审核", "01", 0.18),
        (pts[48], "志愿填报", "02", 0.18),
        (pts[80], "录取", "03", 0.18),
        (pts[106], "四川大学", "04", 0.30),
    ]
    for (px, py), label, num, d in path_nodes:
        add_dot(s, px - d / 2, py - d / 2, d, BLUE if num != "04" else INK)
        if num == "04":
            add_text(s, px - 0.92, py - 0.72, 1.55, 0.38, label, 24, INK, line_spacing=1.0)
            add_text(s, px - 0.04, py + 0.29, 0.55, 0.26, num, 12, BLUE, line_spacing=1.0)
        else:
            add_text(s, px - 0.48, py + 0.28, 1.55, 0.38, label, 24, INK, line_spacing=1.0)
            add_text(s, px - 0.39, py - 0.42, 0.45, 0.26, num, 12, BLUE, line_spacing=1.0)
    add_text(
        s,
        0.84,
        6.42,
        8.9,
        0.5,
        "从报考资格到录取结果，这条路，他走完了。",
        22,
        INK,
        line_spacing=1.5,
        first_indent=True,
    )
    add_source(s, "流程：资格审核 · 志愿填报 · 录取")
    add_page_number(s, 4)
    add_notes(
        s,
        "有了 606 分，还要走完一条具体的录取路：资格审核、志愿填报、录取。专项计划提供了机会，但每一步都要他自己完成。最后，结果落到四川大学。",
    )
    return s


def slide_05(prs):
    s = white_slide(prs)
    add_kicker(s, 5, "等待")
    add_title(s, "等一个确定答案")
    add_text(s, 1.14, 2.12, 5.55, 0.62, "606分之后，", 32, INK, line_spacing=1.1)
    add_text(s, 1.14, 2.95, 6.2, 0.62, "还有一段等待。", 32, BLUE, line_spacing=1.1)
    add_indented_lines(
        s,
        1.16,
        4.12,
        5.75,
        2.05,
        [
            "奶奶说，曾经很心焦。",
            "分数出来，录取还没有落定；",
            "志愿填完，还要等结果。",
            "真正的安心，是录取结果终于到来。",
        ],
        24,
        INK,
    )

    center = (10.52, 3.63)
    rings = [
        (0.92, 0.9, PALE_BLUE, [(18, 155), (178, 318)]),
        (1.38, 2.0, BLUE, [(32, 128), (151, 248), (272, 353)]),
        (1.94, 1.0, PALE_BLUE, [(8, 86), (111, 197), (226, 332)]),
        (2.48, 0.85, PALE_BLUE, [(42, 119), (150, 225), (260, 340)]),
    ]
    for ring_index, (radius, width, color, spans) in enumerate(rings):
        for span_index, (a0, a1) in enumerate(spans):
            count = max(24, int((a1 - a0) * 0.8))
            points = []
            for i in range(count):
                angle = math.radians(a0 + (a1 - a0) * i / (count - 1))
                wobble = 1.0 + 0.065 * math.sin(angle * 3 + ring_index * 0.8)
                wobble += 0.035 * math.sin(angle * 5 - span_index * 0.6)
                points.append(
                    (
                        center[0] + radius * wobble * math.cos(angle),
                        center[1] + radius * wobble * math.sin(angle),
                    )
                )
            add_polyline(s, points, color, width)

    add_text(
        s,
        9.65,
        3.35,
        1.75,
        0.32,
        "等待",
        14,
        BLUE,
        align=PP_ALIGN.CENTER,
        line_spacing=1.0,
    )
    add_text(s, 9.08, 5.82, 3.35, 0.68, "从心焦\n到满意", 24, BLUE, line_spacing=1.5)
    add_source(s, "资料：公开画面与奶奶的引语")
    add_page_number(s, 5)
    add_notes(
        s,
        "等待录取的那些天，家里人并不轻松。奶奶后来用“曾经很心焦”形容那段日子。分数已经出来，结果还没有落定，直到录取结果到来，一家人的心才放下来。",
    )
    return s


def slide_06(prs):
    s = black_slide(prs)
    add_kicker(s, 6, "家人", dark=True)
    add_circle(s, 9.62, 1.16, 5.72, color=BLUE, width=2.8)
    add_title(s, "奶奶的一句话", dark=True)
    add_text(
        s,
        0.98,
        1.96,
        8.25,
        1.1,
        [[("“曾经很心焦，", WHITE, 44)]],
        44,
        WHITE,
        line_spacing=1.2,
    )
    add_text(
        s,
        0.98,
        2.92,
        8.25,
        1.0,
        [[("现在心满意足。”", BLUE, 44)]],
        44,
        BLUE,
        line_spacing=1.2,
    )
    add_rule(s, 10.02, 5.45, 11.98, WHITE, 0.9)
    add_text(s, 9.99, 5.61, 2.0, 0.38, "陈文杰奶奶", 22, WHITE, line_spacing=1.0)
    add_text(
        s,
        1.02,
        5.56,
        6.75,
        1.0,
        "前半句是等待，后半句是落定。",
        24,
        WHITE,
        line_spacing=1.5,
        first_indent=True,
    )
    add_source(s, "引语：公开画面", dark=True)
    add_page_number(s, 6, dark=True)
    add_notes(
        s,
        "奶奶先说心焦，再说心满意足。两个词之间，是一家人等录取结果的时间。现在，结果终于到了。",
    )
    return s


def slide_07(prs):
    s = white_slide(prs)
    add_kicker(s, 7, "心愿")
    add_title(s, "他想把知识带回去")
    add_text(
        s,
        0.98,
        1.76,
        6.55,
        0.88,
        [[("“在大学提升自己，", INK, 40)]],
        40,
        INK,
        line_spacing=1.2,
    )
    add_text(
        s,
        0.98,
        2.60,
        7.1,
        0.88,
        [[("用知识为家乡做贡献。”", BLUE, 40)]],
        40,
        BLUE,
        line_spacing=1.2,
    )
    add_rule(s, 1.02, 3.86, 3.88, INK, 0.9)
    add_text(s, 3.99, 3.72, 2.1, 0.36, "陈文杰", 22, INK, line_spacing=1.0)
    add_indented_lines(
        s,
        1.02,
        4.47,
        6.25,
        1.15,
        ["四川大学是他的下一站，", "家乡是他想回去的地方。"],
        24,
        INK,
    )

    curve = cubic((8.28, 5.45), (9.68, 5.13), (10.58, 2.86), (12.12, 2.08), 90)
    add_polyline(s, curve, BLUE, 2.2)
    add_dot(s, 8.20, 5.37, 0.16, INK)
    add_dot(s, 12.04, 2.00, 0.20, BLUE)
    add_text(s, 8.00, 5.72, 1.1, 0.36, "家乡", 22, INK, line_spacing=1.0)
    add_text(s, 11.18, 1.61, 1.5, 0.36, "大学", 22, BLUE, line_spacing=1.0)
    add_source(s, "引语：公开画面")
    add_page_number(s, 7)
    add_notes(
        s,
        "录取之后，陈文杰想到的不只是自己。他说，要在大学提升自己，用知识为家乡做贡献。四川大学是下一站，家乡是他想回去的地方。",
    )
    return s


def slide_08(prs):
    s = white_slide(prs)
    add_kicker(s, 8, "回响")
    add_title(s, "陈文杰的成绩与心愿，被看见")
    add_text(
        s,
        0.84,
        1.58,
        7.9,
        0.55,
        "屏幕那头的关注，落在他身上。",
        24,
        INK,
        line_spacing=1.5,
        first_indent=True,
    )
    rows = [
        (0.92, 2.34, "01", "他的606分", INK, 2.1),
        (1.62, 3.30, "02", "他的四川大学", BLUE, 4.15),
        (0.92, 4.26, "03", "他想带回家乡的知识", INK, 6.28),
    ]
    for x, y, num, text, color, width in rows:
        add_text(s, x - 0.46, y + 0.02, 0.4, 0.24, num, 12, BLUE, line_spacing=1.0)
        text_width = min(7.45, 9.15 - x)
        add_text(s, x, y, text_width, 0.54, text, 34, color, line_spacing=1.1)
        add_rule(s, x, y + 0.67, x + width, color, 1.0)

    add_vrule(s, 9.42, 2.23, 5.05, color=PALE_BLUE, width=0.8)
    add_text(s, 9.74, 2.21, 2.55, 0.65, "8,476", 40, BLUE, line_spacing=1.0)
    add_text(s, 9.76, 2.85, 1.4, 0.32, "点赞", 18, INK, line_spacing=1.0)
    add_rule(s, 9.75, 3.47, 12.15, PALE_BLUE, 0.7)
    add_text(s, 9.74, 3.74, 2.55, 0.65, "3,286", 36, INK, line_spacing=1.0)
    add_text(s, 9.76, 4.38, 1.4, 0.32, "转发", 18, INK, line_spacing=1.0)
    add_text(
        s,
        0.9,
        5.57,
        8.1,
        1.0,
        "成绩、去向与心愿，让这个故事被看见。",
        24,
        INK,
        line_spacing=1.5,
        first_indent=True,
    )
    add_source(s, "互动数按画面记录：8476 次点赞 / 3286 次转发")
    add_page_number(s, 8)
    add_notes(
        s,
        "这条消息得到 8476 次点赞、3286 次转发。大家看见的，是一个年轻人的成绩、去向和心愿：606 分、四川大学，以及他想带回家乡的知识。",
    )
    return s


def slide_09(prs):
    s = black_slide(prs)
    add_kicker(s, 9, "新的起点", dark=True)
    add_circle(s, 9.02, -1.08, 6.42, color=BLUE, width=2.7)
    add_dot(s, 9.87, 4.10, 0.18, BLUE)
    add_title(s, "出发", dark=True)
    add_text(
        s,
        0.98,
        1.71,
        8.3,
        0.88,
        [[("606分之后，", WHITE, 52)]],
        52,
        WHITE,
        line_spacing=1.1,
    )
    add_text(
        s,
        0.98,
        2.67,
        8.65,
        0.96,
        [[("继续向前。", BLUE, 52)]],
        52,
        BLUE,
        line_spacing=1.1,
    )
    add_rule(s, 1.02, 4.18, 8.55, WHITE, 1.0)
    add_indented_lines(
        s,
        1.04,
        4.55,
        8.15,
        2.0,
        [
            "这张录取结果把陈文杰送进四川大学。",
            "校园、专业、新的生活在前方。",
            "他说，要把知识带回家乡；",
            "这条路，才刚刚开始。",
        ],
        24,
        WHITE,
    )
    add_source(s, "新的路，从这里开始", dark=True)
    add_page_number(s, 9, dark=True)
    add_notes(
        s,
        "这张录取结果把陈文杰送进四川大学。新的校园、新的专业、新的生活还在前面。他说要把知识带回家乡，而这条路才刚刚开始。",
    )
    return s


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    for make_slide in (
        slide_01,
        slide_02,
        slide_03,
        slide_04,
        slide_05,
        slide_06,
        slide_07,
        slide_08,
        slide_09,
    ):
        make_slide(prs)

    for slide in prs.slides:
        strip_style_and_shadow(slide)

    props = prs.core_properties
    props.title = "606分之后｜陈文杰的录取时刻"
    props.subject = "课堂作品汇报"
    props.author = "作品汇报"
    props.keywords = "606分, 陈文杰, 高校专项计划, 四川大学"
    props.comments = "9页；微软雅黑加粗；正文24pt、1.5倍行距、首行缩进2字符。"

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    output = build()
    print(output)
