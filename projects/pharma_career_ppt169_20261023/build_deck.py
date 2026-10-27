# -*- coding: utf-8 -*-
"""Build the 12-page native, editable deck.

Design: Swiss grid, black / white / grey only.
Type: 楷体 for section headings; 华文中宋 for every other glyph,
including numerals and Latin letters.
"""
import math
import pathlib
import re
from functools import lru_cache

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt


P = pathlib.Path(r"C:\Users\camel\Documents\ChatGPT\Hig1\projects\pharma_career_ppt169_20261023")
OUT = P / "_deck_build.pptx"

PX = 9525
W, H = 1280, 720

PAPER = "#F4F4F1"
INK = "#111111"
INK_2 = "#2D2D2B"
MUTED = "#6C6C67"
LIGHT = "#A5A5A0"
RULE = "#D7D7D2"
PALE = "#E8E8E3"
WHITE = "#FFFFFF"

KAI = "楷体"
SONG = "华文中宋"

prs = Presentation()
prs.slide_width = Emu(W * PX)
prs.slide_height = Emu(H * PX)
BLANK = prs.slide_layouts[6]


def _rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


FONT_FILES = {
    KAI: r"C:\Windows\Fonts\simkai.ttf",
    SONG: r"C:\Users\camel\AppData\Local\Microsoft\Windows\Fonts\STZHONGS.TTF",
}

OPEN_PUNCT = "“‘「『（《〈【〔"
CLOSE_PUNCT = "，。；：、！？）》」』】”’…·,.;:!?)]}"

PROTECTED = sorted(
    [
        "制剂研发助理工程师", "药品质量管理", "制剂研发工程师", "岗位能力对照表",
        "仪器分析导论", "基础化学实验", "标准溶液配制", "移液管校准",
        "酸碱滴定", "溶液配制", "原始记录", "开放实验", "药企实习", "药企访谈",
        "制剂工艺", "处方筛选", "工艺优化", "质量研究", "实验技能", "法规知识",
        "数据处理", "数据分析", "实验经验", "规范意识", "持续学习", "实验记录",
        "法规与GMP意识", "数据记录与分析", "GMP意识", "服用时间", "同一成分", "成分", "一次",
        "公开表达", "沟通协作", "药学理论", "实验操作", "复方", "制剂研发",
        "质量研究", "常规型", "研究型", "现实型", "霍兰德测试", "社区医院",
        "降压药", "缓释片", "医药行业", "药学专业", "实验岗", "重复操作",
        "无机化学", "有机化学", "分析化学", "药剂学", "核心课", "综合素养",
        "五个方面", "本科直接做研发", "重复性规范操作", "滴定和配液",
    ],
    key=len,
    reverse=True,
)

TOKEN_RE = re.compile(
    "|".join(
        [re.escape(x) for x in PROTECTED]
        + [
            r"[A-Za-z0-9]+(?:[./\-–—%][A-Za-z0-9]+)*",
            r"[ \u00A0]+",
            r"[\u4e00-\u9fff]",
            r"[^\s]",
        ]
    )
)


@lru_cache(maxsize=None)
def _pil_font(font, size):
    return ImageFont.truetype(FONT_FILES[font], int(round(size)))


def _measure(s, size, font):
    return float(_pil_font(font, size).getlength(s))


def _clean_line(tokens):
    while tokens and tokens[0].isspace():
        tokens.pop(0)
    while tokens and tokens[-1].isspace():
        tokens.pop()
    return "".join(tokens)


def _layout_lines(s, size, width, font):
    tokens = [t for t in TOKEN_RE.findall(s) if t]
    lines = []
    cur = []
    cur_w = 0.0
    for tok in tokens:
        tw = _measure(tok, size, font)
        if tok.isspace():
            if cur and cur_w + tw <= width:
                cur.append(tok)
                cur_w += tw
            continue
        if not cur:
            cur = [tok]
            cur_w = tw
            continue
        if cur_w + tw <= width + 0.01:
            cur.append(tok)
            cur_w += tw
            continue
        if tok and tok[0] in CLOSE_PUNCT:
            cur.append(tok)
            cur_w += tw
            continue
        if cur[-1] and cur[-1][-1] in OPEN_PUNCT and len(cur) > 1:
            moved = cur.pop()
            cur_w -= _measure(moved, size, font)
            line = _clean_line(cur)
            if line:
                lines.append(line)
            cur = [moved, tok]
            cur_w = _measure(moved, size, font) + tw
            continue
        line = _clean_line(cur)
        if line:
            lines.append(line)
        cur = [tok]
        cur_w = tw
    line = _clean_line(cur)
    if line:
        lines.append(line)

    if not lines:
        return [""]

    # Closing punctuation must not begin a line.
    for i in range(1, len(lines)):
        while lines[i] and lines[i][0] in CLOSE_PUNCT:
            lines[i - 1] += lines[i][0]
            lines[i] = lines[i][1:]
    lines = [x for x in lines if x]

    # An opening mark should not be stranded at the end of a line.
    for i in range(len(lines) - 1):
        while lines[i] and lines[i][-1] in OPEN_PUNCT:
            lines[i + 1] = lines[i][-1] + lines[i + 1]
            lines[i] = lines[i][:-1]
        if not lines[i]:
            lines[i] = lines[i + 1]
            lines[i + 1] = ""
    lines = [x for x in lines if x]

    # Avoid a one-character / punctuation-only final line.
    while len(lines) > 1 and len(lines[-1]) <= 1:
        if len(lines[-2]) > 1:
            lines[-1] = lines[-2][-1] + lines[-1]
            lines[-2] = lines[-2][:-1]
        else:
            break
    return lines


def _set_run(run, s, size, font, color, spacing=0, bold=False):
    run.text = s
    run.font.size = Pt(size * 0.75)
    run.font.name = font
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", font)
    if spacing:
        rpr.set("spc", str(int(spacing * 0.75 * 100)))


def _textbox(slide, x, y, w, h, wrap=True):
    tb = slide.shapes.add_textbox(
        Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX))
    )
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return tb


def _draw_line(slide, x, y, w, line, size, font, color, align, spacing, bold):
    lw = _measure(line, size, font)
    if line and line[0] in OPEN_PUNCT:
        left = x - 0.42 * size
    else:
        left = x
    if align == PP_ALIGN.CENTER:
        left = x + (w - lw) / 2.0
    elif align == PP_ALIGN.RIGHT:
        left = x + w - lw
    tb = _textbox(slide, left, y, max(w + size, lw + size * 0.2), size * 1.28, wrap=False)
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = Pt(size * 1.2 * 0.75)
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    r = p.add_run()
    _set_run(r, line, size, font, color, spacing, bold)
    return tb


def text(slide, x, y, w, s, size=20, font=SONG, color=INK, leading=1.62,
         align=PP_ALIGN.LEFT, spacing=0, bold=False):
    lines = _layout_lines(s, size, w, font)
    for i, line in enumerate(lines):
        _draw_line(slide, x, y + i * size * leading, w, line, size, font, color, align, spacing, bold)
    return len(lines)


def paragraphs(slide, x, y, w, items, size=20, font=SONG, color=INK,
               leading=1.62, gap=10, align=PP_ALIGN.LEFT):
    for i, s in enumerate(items):
        lines = _layout_lines(s, size, w, font)
        for j, line in enumerate(lines):
            _draw_line(slide, x, y, w, line, size, font, color, align, 0, False)
            y += size * leading
        if i < len(items) - 1:
            y += gap
    return y


def rect(slide, x, y, w, h, fill=INK, line=None, line_w=1):
    sp = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(int(x * PX)), Emu(int(y * PX)),
        Emu(int(w * PX)), Emu(int(h * PX))
    )
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = _rgb(line)
        sp.line.width = Emu(int(line_w * PX))
    sp.shadow.inherit = False
    return sp


def hline(slide, x1, y, x2, color=RULE, w=1):
    return rect(slide, x1, y, x2 - x1, w, fill=color)


def vline(slide, x, y1, y2, color=RULE, w=1):
    return rect(slide, x, y1, w, y2 - y1, fill=color)


def new_slide(bg=PAPER):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = _rgb(bg)
    return s


def top_meta(slide, page, label, dark=False):
    fg = LIGHT if dark else MUTED
    text(slide, 72, 34, 640, label, 10, SONG, fg, 1.2, spacing=2)
    text(slide, 1088, 34, 120, f"{page:02d} / 12", 10, SONG, fg, 1.2, PP_ALIGN.RIGHT)


def section_head(slide, page, part, title, dark=False):
    fg = WHITE if dark else INK
    rule = "#3A3A38" if dark else INK
    top_meta(slide, page, part, dark)
    text(slide, 72, 60, 900, title, 40, KAI, fg, 1.16)
    hline(slide, 72, 122, 1208, rule, 1)
    if not dark:
        rect(slide, 0, 0, 8, H, INK)


def footer(slide, page, dark=False):
    fg = "#9A9A96" if dark else MUTED
    text(slide, 72, 688, 520, "祝炜程 · 药学 · 大一", 9, SONG, fg, 1.2, spacing=1)
    text(slide, 888, 688, 320, "从药学基础到制药现场", 9, SONG, fg, 1.2, PP_ALIGN.RIGHT, spacing=1)


def card_label(slide, x, y, s, color=MUTED):
    text(slide, x, y, 240, s, 10, SONG, color, 1.2, spacing=1)


def black_strip(slide, x, y, w, h, head, body, head_size=11, body_size=20):
    rect(slide, x, y, w, h, INK)
    text(slide, x + 24, y + 18, w - 48, head, head_size, SONG, "#BDBDB8", 1.2, spacing=1)
    text(slide, x + 24, y + 40, w - 48, body, body_size, SONG, WHITE, 1.45)


# ---------------------------------------------------------------------------
# 01 Cover
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
hline(s, 72, 92, 808, INK, 1)
text(s, 72, 48, 620, "CAREER DEVELOPMENT REPORT / 2026", 9, SONG, MUTED, 1.2, spacing=3)
rect(s, 872, 96, 336, 336, INK)
text(s, 872, 150, 336, "药", 196, KAI, WHITE, 1.0, PP_ALIGN.CENTER)
text(s, 72, 168, 700, "从药学基础", 62, KAI, INK, 1.18)
text(s, 72, 248, 700, "到制药现场", 62, KAI, INK, 1.18)
text(s, 72, 350, 690, "面向制药企业制剂研发助理工程师的成长规划", 20, SONG, MUTED, 1.55)
hline(s, 72, 444, 572, RULE, 1)
text(s, 72, 470, 220, "成长赛道 / 药学 / 大一", 13, SONG, INK_2, 1.4, spacing=1)
text(s, 72, 510, 320, "祝炜程", 34, SONG, INK, 1.2)
rect(s, 0, 592, 1280, 128, INK)
text(s, 72, 620, 500, "PHARMACY / CAREER PLANNING", 9, SONG, "#A9A9A4", 1.2, spacing=3)
text(s, 72, 650, 600, "从一次取药，到一条可执行的成长路径", 16, SONG, WHITE, 1.3)
text(s, 1000, 648, 208, "2026.10 / 12 PAGES", 10, SONG, "#A9A9A4", 1.2, PP_ALIGN.RIGHT, spacing=2)


# ---------------------------------------------------------------------------
# 02 Contents
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
rect(s, 0, 0, 320, H, INK)
text(s, 48, 80, 220, "目录", 54, KAI, WHITE, 1.2)
text(s, 48, 158, 220, "CONTENTS", 10, SONG, "#A9A9A4", 1.2, spacing=4)
text(s, 48, 520, 220, "12", 88, SONG, WHITE, 1.0)
text(s, 48, 624, 220, "PAGES / FULL TEXT", 9, SONG, "#A9A9A4", 1.2, spacing=2)
top_meta(s, 2, "CONTENTS / 10 SECTIONS")
items = [
    ("01", "职业缘起", "03"),
    ("02", "自我认知", "04"),
    ("03", "目标分解", "05"),
    ("04", "差距识别", "06"),
    ("05", "具体行动与成果", "07"),
    ("06", "触发信号", "08"),
    ("07", "调整决策", "09"),
    ("08", "衔接行动", "10"),
    ("09", "阶段性成效", "11"),
    ("10", "总结展望", "12"),
]
for i, (num, title, page) in enumerate(items):
    col = i % 2
    row = i // 2
    x = 376 + col * 424
    y = 96 + row * 106
    text(s, x, y, 34, num, 10, SONG, MUTED, 1.2, spacing=1)
    text(s, x + 48, y - 8, 320, title, 24, SONG, INK, 1.3)
    text(s, x + 332, y, 44, page, 11, SONG, MUTED, 1.2, PP_ALIGN.RIGHT)
    hline(s, x, y + 48, x + 376, RULE, 1)
text(s, 888, 688, 320, "从药学基础到制药现场", 9, SONG, MUTED, 1.2, PP_ALIGN.RIGHT, spacing=1)


# ---------------------------------------------------------------------------
# 03 1.1 职业缘起
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
section_head(s, 3, "PART 01 / 一、职业目标的设定过程", "1.1  职业缘起")
p1 = [
    "我选择制药方向，起因是陪奶奶去社区医院取降压药。药师把每种药的服用时间写在盒子上，又提醒不能自己掰开缓释片。那时我才知道，同一成分做成不同剂型，患者用起来差别很大。",
    "我家在县城，父母做普通工作，家里没有医药行业的人脉。填志愿时，家里最关心的是毕业后能不能顺利就业。我查了药学专业的去向，发现制药企业有一条从实验岗做起的路径，门槛相对清楚，也和我坐得住、愿意做重复操作的性子合得来。",
    "我把短期目标定为制药企业制剂研发助理工程师，中期争取成为制剂研发工程师；备选方向是药品质量管理QA/QC。",
]
paragraphs(s, 72, 154, 652, p1, 20, SONG, INK, 1.66, 12)

card_label(s, 768, 154, "从一次取药开始 / ORIGIN")
steps = [
    ("01", "陪奶奶去社区医院取降压药"),
    ("02", "药师把每种药的服用时间写在盒子上"),
    ("03", "提醒不能自己掰开缓释片"),
    ("04", "同一成分，不同剂型，患者用起来差别很大"),
]
for i, (num, body) in enumerate(steps):
    yy = 190 + i * 88
    rect(s, 768, yy, 34, 34, INK)
    text(s, 768, yy + 7, 34, num, 10, SONG, WHITE, 1.0, PP_ALIGN.CENTER)
    text(s, 820, yy - 2, 388, body, 16, SONG, INK, 1.5)
    if i < 3:
        hline(s, 768, yy + 56, 1208, RULE, 1)
rect(s, 768, 556, 440, 104, INK)
text(s, 792, 572, 392, "短期目标", 10, SONG, "#BDBDB8", 1.2, spacing=1)
text(s, 792, 594, 392, "制剂研发助理工程师", 20, SONG, WHITE, 1.45)
text(s, 792, 640, 392, "中期：制剂研发工程师 / 备选：QA/QC", 9, SONG, "#C4C4BE", 1.2)
footer(s, 3)


# ---------------------------------------------------------------------------
# 04 1.2 自我认知
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
section_head(s, 4, "PART 01 / 一、职业目标的设定过程", "1.2  自我认知")
p2 = [
    "我做过MBTI和霍兰德测试，结果是ISTJ，霍兰德代码CIR，以常规型为主，也带研究型和现实型。",
    "我的优势是做事细、耐心足，实验记录愿意写全；短板是公开表达容易紧张，数据分析基础弱，实验经验也少。",
    "图表1显示，我的细致和规范意识与制剂研发所需的严谨、规范记录和持续学习较匹配；最大差距在实验技能、法规知识和数据处理。",
]
paragraphs(s, 72, 154, 668, p2, 22, SONG, INK, 1.7, 12)
rect(s, 72, 486, 640, 116, INK)
text(s, 96, 504, 592, "图表 1 的结论", 10, SONG, "#BDBDB8", 1.2, spacing=1)
text(s, 96, 528, 592, "我的细致和规范意识与制剂研发所需的严谨、", 18, SONG, WHITE, 1.45)
text(s, 96, 554, 592, "规范记录和持续学习较匹配。", 18, SONG, WHITE, 1.45)

card_label(s, 768, 154, "MBTI / HOLLAND")
text(s, 768, 182, 420, "ISTJ", 72, SONG, INK, 1.05, spacing=1)
text(s, 768, 264, 420, "CIR", 72, SONG, INK, 1.05, spacing=2)
text(s, 768, 350, 420, "常规型为主 / 研究型 / 现实型", 14, SONG, MUTED, 1.4)
hline(s, 768, 394, 1208, RULE, 1)
card_label(s, 768, 410, "优势")
text(s, 768, 432, 420, "做事细 / 耐心足 / 实验记录愿意写全", 17, SONG, INK, 1.55)
card_label(s, 768, 506, "短板")
text(s, 768, 528, 420, "公开表达容易紧张 / 数据分析基础弱 / 实验经验也少", 17, SONG, INK, 1.55)
footer(s, 4)


# ---------------------------------------------------------------------------
# 05 1.4 目标分解
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
section_head(s, 5, "PART 01 / 一、职业目标的设定过程", "1.4  目标分解")
p3 = [
    "短期从大一到毕业，学好无机化学、有机化学和分析化学，再学药剂学等核心课；实验上练熟移液、滴定和溶液配制，原始记录也要写完整；大一下学期争取进课题组做开放实验，大二暑假找一份药企实习；毕业前完成一次药企访谈，并更新岗位能力对照表。",
    "中期是毕业后1到3年，进入制药企业做制剂研发助理工程师或QA/QC，熟悉制剂工艺、GMP和验证。",
    "长期是3到5年，能独立承担处方筛选、工艺优化或质量研究中的一项。",
]
paragraphs(s, 72, 152, 1056, p3, 21, SONG, INK, 1.68, 12)

cols = [
    ("短期 / 大一到毕业", ["核心课、移液、滴定、溶液配制", "开放实验、药企实习、药企访谈"]),
    ("中期 / 毕业后 1–3 年", ["制剂研发助理工程师或QA/QC", "制剂工艺、GMP、验证"]),
    ("长期 / 3–5 年", ["独立承担处方筛选、工艺优化或质量研究中的一项"]),
]
for i, (head, body_lines) in enumerate(cols):
    x = 72 + i * 392
    rect(s, x, 468, 6, 6, INK)
    text(s, x + 18, 458, 340, head, 18, KAI, INK, 1.35)
    hline(s, x, 496, x + 340, INK, 1)
    paragraphs(s, x, 514, 340, body_lines, 14, SONG, INK_2, 1.65, 4)
footer(s, 5)


# ---------------------------------------------------------------------------
# 06 2.1 差距识别
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
section_head(s, 6, "PART 02 / 二、实现职业目标的具体行动与成效", "2.1  差距识别")
p4 = [
    "图表3把岗位要求拆成5项能力：",
    "药学理论，实验操作，数据记录与分析，",
    "法规与GMP意识，沟通协作。",
]
paragraphs(s, 72, 164, 620, p4, 26, SONG, INK, 1.7, 0)
text(s, 72, 334, 620, "对照下来，我差距最大的是实验操作和法规意识。", 26, SONG, INK, 1.7)
black_strip(s, 72, 454, 548, 114, "最大差距", "实验操作 / 法规意识", 10, 28)
text(s, 72, 614, 548, "自评仅用于排序，不是测评分数。", 10, SONG, MUTED, 1.3)

cap = [
    ("药学理论", 3, ""),
    ("实验操作", 1, "P1"),
    ("数据记录与分析", 2, ""),
    ("法规与GMP意识", 1, "P1"),
    ("沟通协作", 2, ""),
]
for i, (label, score, tag) in enumerate(cap):
    yy = 168 + i * 88
    text(s, 700, yy + 4, 190, label, 14, SONG, INK, 1.3)
    for j in range(5):
        fill = INK if j < score else PALE
        rect(s, 900 + j * 38, yy, 30, 28, fill)
    text(s, 1110, yy + 3, 44, f"{score}/5", 11, SONG, MUTED, 1.2)
    if tag:
        rect(s, 1162, yy, 34, 28, INK)
        text(s, 1162, yy + 7, 34, tag, 9, SONG, WHITE, 1.0, PP_ALIGN.CENTER)
    hline(s, 700, yy + 48, 1208, RULE, 1)
text(s, 700, 634, 508, "五项能力 = 岗位要求的最小检查表", 10, SONG, MUTED, 1.2)
footer(s, 6)


# ---------------------------------------------------------------------------
# 07 2.2 具体行动与成果
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
section_head(s, 7, "PART 02 / 二、实现职业目标的具体行动与成效", "2.2  具体行动与成果")
p5 = [
    "针对差距，我的行动覆盖从课程学习到综合素养的五个方面。",
    "课程上，我把无机化学当作重点。",
    "实验上，大一上学期完成多次基础化学实验，",
    "内容包括酸碱滴定、移液管校准和标准溶液配制。",
]
paragraphs(s, 72, 164, 660, p5, 23, SONG, INK, 1.72, 10)
black_strip(s, 72, 474, 604, 96, "动作发生在大一上学期", "多次基础化学实验 / 三项基础内容", 10, 20)

card_label(s, 748, 158, "课程 / COURSES")
text(s, 748, 182, 460, "无机化学", 40, SONG, INK, 1.2)
text(s, 748, 240, 460, "有机化学 / 分析化学 / 药剂学等核心课", 17, SONG, INK_2, 1.5)
hline(s, 748, 300, 1208, INK, 1)
card_label(s, 748, 320, "实验 / LAB")
text(s, 748, 344, 460, "移液 / 滴定", 30, SONG, INK, 1.35)
text(s, 748, 390, 460, "溶液配制", 30, SONG, INK, 1.35)
hline(s, 748, 454, 1208, RULE, 1)
card_label(s, 748, 474, "记录 / RECORD")
text(s, 748, 498, 460, "原始记录写完整", 28, SONG, INK, 1.35)
text(s, 748, 552, 460, "课程学习 · 实验操作 · 综合素养", 16, SONG, MUTED, 1.4)
footer(s, 7)


# ---------------------------------------------------------------------------
# 08 3.1 触发信号
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
section_head(s, 8, "PART 03 / 三、职业目标及行动的动态调整", "3.1  触发信号")
p6 = [
    "本科直接做研发的机会不多。另一方面，我在重复性规范操作上反而更稳，做滴定和配液时愿意一遍遍核对，出错少。",
    "两件事放在一起，我开始怀疑原先“直接冲制剂研发”的路径是不是太窄。",
]
paragraphs(s, 72, 158, 1056, p6, 25, SONG, INK, 1.72, 16)

signals = [
    ("01", "药企参观", "招聘要求里，硕士和有项目经历的人更受青睐。"),
    ("02", "校友访谈", "本科直接做研发的机会不多。"),
    ("03", "自我观察", "重复性规范操作更稳，滴定和配液愿意一遍遍核对，出错少。"),
]
for i, (num, head, body) in enumerate(signals):
    x = 72 + i * 384
    rect(s, x, 374, 34, 34, INK)
    text(s, x, 381, 34, num, 10, SONG, WHITE, 1.0, PP_ALIGN.CENTER)
    text(s, x + 48, 374, 300, head, 20, SONG, INK, 1.35)
    text(s, x, 424, 356, body, 15, SONG, INK_2, 1.7)
    hline(s, x, 524, x + 356, RULE, 1)
black_strip(s, 72, 560, 1136, 104, "问题被重新提出", "直接冲制剂研发，路径是不是太窄？", 10, 28)
footer(s, 8)


# ---------------------------------------------------------------------------
# 09 3.2 调整决策
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
section_head(s, 9, "PART 03 / 三、职业目标及行动的动态调整", "3.2  调整决策")
p7 = [
    "调整后的短期目标是把实验、法规和数据基础打牢，大二上学期再在制剂研发与QA/QC之间做二次选择；备选方向从“万一”变成“并行准备”。",
    "不变的是进入制药行业、守住规范与质量这条主线。家里没有行业人脉，考研又要多花两三年，开销也得算。我不想把它当成逃避就业的借口，所以会先看岗位要求，再决定要不要考。",
]
paragraphs(s, 72, 152, 1090, p7, 21, SONG, INK, 1.68, 12)

rect(s, 72, 392, 462, 202, PALE)
card_label(s, 96, 414, "原路径 / BEFORE")
text(s, 96, 440, 402, "直接冲制剂研发", 28, SONG, INK, 1.35)
text(s, 96, 500, 402, "备选方向只当“万一”", 15, SONG, MUTED, 1.5)

rect(s, 566, 392, 642, 202, INK)
card_label(s, 590, 414, "调整后 / AFTER", "#B9B9B4")
text(s, 590, 440, 594, "先打牢实验、法规和数据基础", 24, SONG, WHITE, 1.45)
text(s, 590, 500, 594, "大二上：制剂研发 / QA/QC 二次选择", 15, SONG, "#D6D6D1", 1.5)
text(s, 590, 528, 594, "备选：从“万一”变成“并行准备”", 15, SONG, "#D6D6D1", 1.5)
text(s, 590, 556, 594, "考研：先看岗位要求，再决定要不要考", 15, SONG, "#D6D6D1", 1.5)
text(s, 72, 628, 1136, "不变的主线：进入制药行业，守住规范与质量。", 18, SONG, INK, 1.4)
footer(s, 9)


# ---------------------------------------------------------------------------
# 10 3.3 衔接行动
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
section_head(s, 10, "PART 03 / 三、职业目标及行动的动态调整", "3.3  衔接行动")
p8 = [
    "已有课程和实验成果仍然是基础。",
    "新增的行动是把选课表里的一门选修换成了仪器分析导论。",
    "调整只是把路走稳一点，原有课程和实验成果都还在。",
]
paragraphs(s, 72, 158, 684, p8, 24, SONG, INK, 1.74, 16)

rect(s, 808, 158, 400, 176, PALE)
card_label(s, 832, 180, "已有 / KEEP")
text(s, 832, 204, 352, "课程 / 实验成果", 24, SONG, INK, 1.4)
text(s, 832, 254, 352, "仍然是基础", 15, SONG, MUTED, 1.4)

rect(s, 808, 358, 400, 176, INK)
card_label(s, 832, 380, "新增 / ADD", "#B9B9B4")
text(s, 832, 404, 352, "仪器分析导论", 28, SONG, WHITE, 1.4)
text(s, 832, 456, 352, "一门选修的替换", 15, SONG, "#D6D6D1", 1.4)

black_strip(s, 72, 558, 1136, 102, "调整的边界", "只是把路走稳一点，原有课程和实验成果都还在。", 10, 24)
footer(s, 10)


# ---------------------------------------------------------------------------
# 11 阶段性成效
# ---------------------------------------------------------------------------
s = new_slide(PAPER)
section_head(s, 11, "PART 09 / 阶段回看", "阶段性成效")

cards = [
    ("多次", "基础化学实验", ["酸碱滴定、移液管校准、标准溶液配制"], "已完成 · 大一上"),
    ("5项", "岗位能力拆解", ["药学理论、实验操作、数据记录与分析", "法规与GMP意识、沟通协作"], "已明确 · 最大差距两项"),
    ("1门", "选修调整", ["仪器分析导论"], "已调整 · 衔接行动"),
]
for i, (big, head, body_lines, tag) in enumerate(cards):
    x = 72 + i * 392
    rect(s, x, 196, 336, 222, WHITE)
    rect(s, x, 196, 336, 6, INK)
    text(s, x + 22, 218, 160, big, 58, SONG, INK, 1.0)
    text(s, x + 22, 292, 292, head, 18, SONG, INK, 1.35)
    paragraphs(s, x + 22, 324, 292, body_lines, 12, SONG, INK_2, 1.9, 4)
    text(s, x + 22, 390, 292, tag, 10, SONG, MUTED, 1.2)

hline(s, 72, 500, 1208, INK, 1)
milestones = [
    ("大一下", "开放实验"),
    ("大二暑假", "药企实习"),
    ("毕业前", "药企访谈"),
    ("毕业后 1–3 年", "助理工程师 / QA/QC"),
    ("3–5 年", "独立承担一项"),
]
for i, (when, what) in enumerate(milestones):
    x = 72 + i * 228
    rect(s, x, 494, 10, 10, INK)
    text(s, x, 518, 200, when, 11, SONG, MUTED, 1.3)
    text(s, x, 540, 200, what, 15, SONG, INK, 1.45)
footer(s, 11)


# ---------------------------------------------------------------------------
# 12 结尾 / 总结展望
# ---------------------------------------------------------------------------
s = new_slide(INK)
section_head(s, 12, "PART 10 / 结尾", "总结展望", dark=True)
p9 = [
    "制药对我而言，是用规范、数据和实验能力，把安全有效的药品做出来。",
    "我能依靠的主要是课程、实验和一次次访谈得到的信息。",
]
paragraphs(s, 72, 168, 910, p9, 24, SONG, "#E8E8E3", 1.72, 16)
hline(s, 72, 356, 1208, "#3A3A38", 1)
text(s, 72, 390, 1040, "大一的目标是先成为“可被培养的制药人”，", 34, SONG, WHITE, 1.45)
text(s, 72, 448, 1040, "再成为“能解决问题的工程师”。", 34, SONG, "#BDBDB8", 1.45)
footer(s, 12, dark=True)


# ---------------------------------------------------------------------------
# Notes, flatness, save
# ---------------------------------------------------------------------------
notes = {
    1: "封面。题目与身份信息。",
    2: "目录。全文按小节展开，共 12 页。",
    3: "1.1 职业缘起。从陪奶奶取降压药讲到剂型差异，再落到实验岗路径。",
    4: "1.2 自我认知。MBTI 与霍兰德只作为自我观察入口。",
    5: "1.4 目标分解。短期、中期、长期三个时间尺度。",
    6: "2.1 差距识别。五项能力中，实验操作与法规意识差距最大。",
    7: "2.2 具体行动与成果。课程重点与大一上的基础实验。",
    8: "3.1 触发信号。参观、访谈与自我观察共同触发路径反思。",
    9: "3.2 调整决策。短期打基础，大二上二次选择，备选并行准备。",
    10: "3.3 衔接行动。只替换一门选修，原有成果保留。",
    11: "阶段性成效。只汇总前文出现的证据，计划中的事项单独标注。",
    12: "结尾。职业价值与两阶段目标。",
}
for idx, slide in enumerate(prs.slides, start=1):
    slide.notes_slide.notes_text_frame.text = notes[idx]


for slide in prs.slides:
    for shape in slide.shapes:
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)
        sp_pr = shape._element.find(qn("p:spPr"))
        if sp_pr is not None and sp_pr.find(qn("a:effectLst")) is None:
            etree.SubElement(sp_pr, qn("a:effectLst"))

prs.save(str(OUT))
print(f"saved {len(prs.slides)} slides -> {OUT}")

