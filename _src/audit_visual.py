# -*- coding: utf-8 -*-
"""Audit the three-colour, contour and transition-block requirements."""
from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation

MORANDI = {"B56F63", "71877D", "B39A78", "D49286", "8FA79C", "D6BE9A"}

path = Path(sys.argv[1] if len(sys.argv) > 1 else
            r"projects\zhiwen_teaching_ppt169_20261024\蛭文_教学读本_95页.pptx")
prs = Presentation(path)
bad_colour, bad_contour, bad_transition = [], [], []
summary = []

for idx, slide in enumerate(prs.slides, 1):
    names = []
    fills = set()
    contour_lines = []
    for shape in slide.shapes:
        names.append(shape.name)
        if shape.name == "bgcurve_contour":
            try:
                contour_lines.append(round(shape.line.width.pt, 2))
            except Exception:
                pass
        try:
            if shape.fill.type is not None and shape.fill.type == 1:
                fills.add(str(shape.fill.fore_color.rgb).upper())
        except Exception:
            pass
    triad = sorted(fills & MORANDI)
    if len(triad) < 3:
        bad_colour.append({"slide": idx, "morandi": triad, "all_fills": sorted(fills)})
    if not any(n == "bgcurve_contour" for n in names):
        bad_contour.append(idx)
    if not (any(n == "bgtransition_mid" for n in names)
            and any(n == "bgtransition_edge" for n in names)):
        bad_transition.append(idx)
    summary.append({
        "slide": idx,
        "triad": triad,
        "contours": sum(n == "bgcurve_contour" for n in names),
        "contour_pt": contour_lines,
    })

out = {
    "file": str(path),
    "slides": len(prs.slides),
    "bad_colour": bad_colour,
    "bad_contour": bad_contour,
    "bad_transition": bad_transition,
    "n_bad": len(bad_colour) + len(bad_contour) + len(bad_transition),
    "details": summary,
}
Path(r"_src\visualaudit.json").write_text(
    json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps({k: out[k] for k in
                  ("file", "slides", "n_bad", "bad_colour",
                   "bad_contour", "bad_transition")},
                 ensure_ascii=False, indent=1))
