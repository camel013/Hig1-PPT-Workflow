import sys
from pptx import Presentation
from pptx.util import Emu
src=sys.argv[1]
pages=[int(x) for x in sys.argv[2].split(',')]
p=Presentation(src)
for n in pages:
    s=p.slides[n-1]
    print(f'\n===== PAGE {n:02d} =====')
    shapes=[]
    for sh in s.shapes:
        if not getattr(sh,'has_text_frame',False): continue
        txt=' | '.join(x.strip().replace('\n',' / ') for x in sh.text_frame.text.split('\n') if x.strip())
        if not txt: continue
        sizes=[]
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if r.font.size: sizes.append(r.font.size.pt)
        fs=max(sizes) if sizes else 0
        shapes.append((round(sh.top/12700),round(sh.left/12700),round(sh.width/12700),round(sh.height/12700),fs,txt))
    for sh in sorted(shapes): print(f'y{sh[0]:3} x{sh[1]:4} w{sh[2]:4} h{sh[3]:3} s{sh[4]:4.1f} | {sh[5][:240]}')
