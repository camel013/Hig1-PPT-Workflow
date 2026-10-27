import re
import sys
from pathlib import Path

from pptx import Presentation


FOOTER = "《蛭文》教学读本 / 文本细读与写作"
SOURCE = re.compile(r"^《[^》]{1,30}》(?:\s*/\s*[^/]{1,24}|开篇|节选|跋|（节选）)?$")
DASH_SOURCE = re.compile(r"^——\s*[^—]{1,16}$")
EMU = 12700


def paragraphs(shape):
    out = []
    for para in shape.text_frame.paragraphs:
        text = "".join(run.text for run in para.runs).strip()
        if text:
            out.append(text)
    return out


def font_size(shape):
    sizes = []
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size:
                sizes.append(run.font.size.pt)
    return max(sizes) if sizes else 0


def info(shape):
    return {
        "shape": shape,
        "text": "\n".join(paragraphs(shape)),
        "x": shape.left / EMU,
        "y": shape.top / EMU,
        "w": shape.width / EMU,
        "h": shape.height / EMU,
        "size": font_size(shape),
    }


def table_text(shape):
    rows = []
    for row in shape.table.rows:
        cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def filled_rects(slide):
    rects = []
    for shape in slide.shapes:
        if shape.has_table:
            continue
        if shape.has_text_frame and shape.text_frame.text.strip():
            continue
        if shape.shape_type != 1:
            continue
        try:
            fill = str(shape.fill.fore_color.rgb)
        except Exception:
            continue
        x, y = shape.left / EMU, shape.top / EMU
        w, h = shape.width / EMU, shape.height / EMU
        if 40 <= w <= 900 and 12 <= h <= 500 and w * h < 420000:
            rects.append((x, y, w, h, fill))
    return rects


def containing_rect(rects, item):
    cx = item["x"] + item["w"] / 2
    cy = item["y"] + item["h"] / 2
    inside = [r for r in rects if r[0] <= cx <= r[0] + r[2] and r[1] <= cy <= r[1] + r[3]]
    inside.sort(key=lambda r: r[2] * r[3])
    return inside[0] if inside else None


def is_source(text):
    return len(text) <= 42 and (bool(SOURCE.fullmatch(text)) or bool(DASH_SOURCE.fullmatch(text)))


def collect(slide, page):
    if page == 1:
        return (
            "蛭文\n\n一本假课本的读法\n\n"
            "从戏仿到见证：语言如何成为证词\n\n"
            "高中语文 · 文本细读与写作\n\n"
            "4 个单元 · 24 篇课文 · 3 首古诗词 · 序与跋"
        )

    items = []
    tables = []
    vertical = []
    for shape in slide.shapes:
        if shape.has_table:
            text = table_text(shape)
            if text:
                tables.append({"text": text, "y": shape.top / EMU, "x": shape.left / EMU})
            continue
        if not shape.has_text_frame:
            continue
        item = info(shape)
        if not item["text"]:
            continue
        if item["y"] < 40 or item["text"] == FOOTER:
            continue
        if re.fullmatch(rf"{page:02d}", item["text"]) and item["y"] > 480:
            continue
        if re.fullmatch(r"\d{1,3} / 95", item["text"]):
            continue
        if re.fullmatch(r"\d{1,2}", item["text"]) and item["h"] > 50:
            continue
        if item["w"] <= 70 and item["h"] <= 40 and len(item["text"]) <= 2:
            vertical.append(item)
            continue
        items.append(item)

    source_items = [item for item in items if is_source(item["text"])]
    rects = filled_rects(slide)
    remove = set()

    for src in source_items:
        remove.add(id(src["shape"]))
        rect = containing_rect(rects, src)
        if rect and len(source_items) <= 2:
            for item in items:
                if item is src:
                    continue
                cx = item["x"] + item["w"] / 2
                cy = item["y"] + item["h"] / 2
                if (rect[0] <= cx <= rect[0] + rect[2]
                        and rect[1] <= cy <= rect[1] + rect[3]
                        and item["y"] < src["y"] - 2
                        and item["size"] >= 9):
                    remove.add(id(item["shape"]))
        else:
            nearest = None
            for item in items:
                if item is src or item["y"] >= src["y"]:
                    continue
                gap = src["y"] - (item["y"] + item["h"])
                if gap < 0:
                    continue
                if not (item["x"] < src["x"] + src["w"] and item["x"] + item["w"] > src["x"]):
                    continue
                if 0 <= gap <= 45 and item["size"] <= 9.5 and len(item["text"]) >= 8:
                    if nearest is None or gap < nearest[0]:
                        nearest = (gap, item)
            if nearest:
                remove.add(id(nearest[1]["shape"]))

    if page == 4:
        for item in items:
            if item["size"] >= 20 and 90 <= item["y"] <= 260:
                remove.add(id(item["shape"]))

    if page == 21:
        for item in items:
            if 12 <= item["size"] <= 13.5 and item["y"] < 460:
                remove.add(id(item["shape"]))

    remaining = [item for item in items if id(item["shape"]) not in remove]

    title = None
    for item in sorted(remaining, key=lambda x: (x["y"], x["x"])):
        if item["y"] < 110 and item["size"] >= 20:
            title = item["text"]
            remove.add(id(item["shape"]))
            break

    body = [item for item in remaining if id(item["shape"]) not in remove]
    body.sort(key=lambda x: (x["y"], x["x"]))
    lines = []
    if title:
        lines.append(title)
    for item in body:
        if not lines or lines[-1] != item["text"]:
            lines.append(item["text"])
    for table in sorted(tables, key=lambda x: (x["y"], x["x"])):
        lines.append(table["text"])
    return "\n\n".join(lines)


def main():
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    pres = Presentation(str(src))
    chunks = []
    for i, slide in enumerate(pres.slides, 1):
        body = collect(slide, i)
        chunks.append(f"## 第{i:02d}页\n\n{body}")
    header = (
        "# 《蛭文》教学读本 · 讲解文字稿\n\n"
        f"共 {len(pres.slides)} 页。只保留讲解、分析与教学提示，不含引文原文、诗作原文与出处标注。\n"
    )
    dst.write_text(header + "\n\n---\n\n" + "\n\n---\n\n".join(chunks) + "\n", encoding="utf-8")
    print(f"wrote {dst} ({dst.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
