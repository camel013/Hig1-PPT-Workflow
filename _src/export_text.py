import re
import sys
from pathlib import Path

from pptx import Presentation


FOOTER = "《蛭文》教学读本 / 文本细读与写作"
EMU = 12700


def frame_text(shape):
    lines = []
    for para in shape.text_frame.paragraphs:
        text = "".join(run.text for run in para.runs).strip()
        if text:
            lines.append(text)
    return "\n".join(lines)


def table_text(shape):
    rows = []
    for row in shape.table.rows:
        cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def collect(slide, page):
    normal = []
    columns = {}
    for shape in slide.shapes:
        if shape.has_table:
            text = table_text(shape)
            if text:
                normal.append((shape.top / EMU, shape.left / EMU, text))
            continue
        if not shape.has_text_frame:
            continue
        text = frame_text(shape)
        if not text:
            continue
        top, left = shape.top / EMU, shape.left / EMU
        width, height = shape.width / EMU, shape.height / EMU

        if top < 40:
            continue
        if text == FOOTER:
            continue
        if re.fullmatch(rf"{page:02d}", text) and top > 480:
            continue
        if re.fullmatch(r"\d{1,3} / 95", text):
            continue
        if re.fullmatch(r"\d{1,2}", text) and height > 50:
            continue

        if width <= 50 and height <= 35 and len(text) <= 2:
            key = round(left / 4) * 4
            columns.setdefault(key, []).append((top, text))
            continue
        normal.append((top, left, text))

    column_lines = []
    for x, chars in columns.items():
        chars.sort()
        if len(chars) >= 2:
            column_lines.append({
                "x": x,
                "top": chars[0][0],
                "text": "".join(ch for _, ch in chars),
            })

    used = set()
    for i, line in enumerate(column_lines):
        if i in used:
            continue
        cluster = [j for j, other in enumerate(column_lines)
                   if abs(other["top"] - line["top"]) <= 24]
        if len(cluster) >= 3:
            group = [column_lines[j] for j in cluster]
            group.sort(key=lambda item: item["x"], reverse=True)
            normal.append((line["top"], max(item["x"] for item in group),
                           "\n".join(item["text"] for item in group)))
            used.update(cluster)
        else:
            normal.append((line["top"], line["x"], line["text"]))
            used.add(i)

    normal.sort(key=lambda item: (item[0], item[1]))
    out = []
    for _, _, text in normal:
        if not out or out[-1] != text:
            out.append(text)
    return "\n\n".join(out)


def main():
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    pres = Presentation(str(src))
    chunks = []
    for i, slide in enumerate(pres.slides, 1):
        body = collect(slide, i)
        chunks.append(f"## 第{i:02d}页\n\n{body}")
    header = (
        "# 《蛭文》教学读本 · 文字稿\n\n"
        f"共 {len(pres.slides)} 页。按页面阅读顺序整理，仅保留页面文字，不含版式与配色说明。\n"
    )
    dst.write_text(header + "\n\n---\n\n" + "\n\n---\n\n".join(chunks) + "\n", encoding="utf-8")
    print(f"wrote {dst} ({dst.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
