from __future__ import annotations

import hashlib
import html
import json
import subprocess
import tempfile
import zipfile
from collections import Counter
from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt


BASE = Path(__file__).resolve().parent.parent
WORKFLOW = BASE / "workflow"
STYLE_PATH = WORKFLOW / "swiss_duoyuhua_v1.json"
OUTPUT = BASE / "deliverables"
DONOR = Path(
    r"C:\Users\camel\Documents\ChatGPT\Hig1\projects"
    r"\duoyuhua_ppt169_20261021\多余的话_瑞士国际主义版式.pptx"
)
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
REL_NS = {
    "r": "http://schemas.openxmlformats.org/package/2006/relationships",
}
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def text(value, x, y, type_role, width, color="ink", font_role="body",
         align="left", semantic_role=None, name=None):
    return {
        "kind": "text",
        "text": value,
        "x": x,
        "y": y,
        "w": width,
        "h": 110,
        "type_role": type_role,
        "role": semantic_role or type_role,
        "color": color,
        "font_role": font_role,
        "align": align,
        "name": name or f"text_{abs(hash((value, x, y))) & 0xFFFFFF:06x}",
    }


def rect(x, y, w, h, fill=None, stroke=None, sw=0, role="shape", name=None):
    return {
        "kind": "rect",
        "x": x,
        "y": y,
        "w": w,
        "h": h,
        "fill": fill,
        "stroke": stroke,
        "stroke_width_px": sw,
        "role": role,
        "name": name or f"shape_{abs(hash((x, y, w, h))) & 0xFFFFFF:06x}",
    }


def oval(x, y, w, h, fill=None, stroke=None, sw=0, role="shape", name=None):
    return {
        "kind": "oval",
        "x": x,
        "y": y,
        "w": w,
        "h": h,
        "fill": fill,
        "stroke": stroke,
        "stroke_width_px": sw,
        "role": role,
        "name": name or f"shape_{abs(hash((x, y, w, h))) & 0xFFFFFF:06x}",
    }


def frame(page_count, page, section, right, footer):
    return [
        text(section, 80, 62, "micro", 1040, color="gray", semantic_role="chrome"),
        text(right, 920, 62, "micro", 280, color="gray", align="right",
             semantic_role="chrome"),
        rect(80, 96, 1120, 1, fill="hair", role="chrome_rule"),
        rect(80, 664, 1120, 1, fill="hair", role="chrome_rule"),
        text(footer, 80, 680, "micro", 1000, color="gray", semantic_role="chrome"),
        text(f"{page:02d} / {page_count:02d}", 1040, 680, "micro", 160,
             color="gray", align="right", semantic_role="chrome"),
    ]


def meeting_pages():
    one = [
        rect(0, 0, 36, 720, fill="terracotta", role="chrome"),
        rect(1244, 0, 36, 720, fill="blue", role="chrome"),
        oval(907, -245, 518, 518, fill="blue", role="hero_arc"),
        oval(925, -215, 470, 470, stroke="gray", sw=1),
        oval(985, -171, 360, 360, stroke="gray", sw=1),
        oval(1045, -127, 250, 250, stroke="gray", sw=1),
        text("TEAM EFFICIENCY / BASELINE", 80, 62, "micro", 1000,
             color="gray", semantic_role="chrome"),
        text("2026.10", 900, 62, "micro", 300, color="gray", align="right",
             semantic_role="chrome"),
        text("MEETING / COST / ACTION", 80, 154, "micro", 760,
             color="gray", semantic_role="kicker"),
        text("减少无效", 80, 218, "hero", 760, color="ink", font_role="title",
             semantic_role="hero"),
        text("会议", 80, 322, "hero", 760, color="ink", font_role="title",
             semantic_role="hero"),
        text("4 页最小行动基线", 80, 484, "heading", 760, color="gray"),
        text("MEETING COST", 884, 246, "micro", 300, color="inverse"),
        text("TIME", 884, 326, "label", 300, color="inverse"),
        text("PEOPLE", 884, 380, "label", 300, color="inverse"),
        text("OUTPUT", 884, 434, "label", 300, color="inverse"),
        text("团队效率 / 2026 年 10 月", 80, 680, "micro", 1000,
             color="gray", semantic_role="chrome"),
        text("01 / 04", 1040, 680, "micro", 160, color="gray",
             align="right", semantic_role="chrome"),
    ]
    two = frame(4, 2, "02 / MEETING COST", "时间 / 人数 / 产出",
                "减少无效会议 / 会议成本")
    two += [
        text("会议成本", 80, 126, "title", 1000, color="ink", font_role="title"),
        text("一场会议的成本，", 80, 215, "lead", 520, color="ink", font_role="title"),
        text("远不止一小时", 80, 277, "lead", 520, color="ink", font_role="title"),
        rect(80, 374, 440, 236, fill="terracotta", role="data_block"),
        rect(80, 374, 6, 236, fill="ink"),
        text("直接成本 / 6 人 × 30 分钟", 116, 399, "label", 360,
             color="inverse"),
        text("6 × 30", 116, 437, "metric", 360, color="inverse"),
        text("= 3 人小时", 116, 518, "heading", 360, color="inverse"),
        text("半小时的共识，会占用多人的工作上下文。", 116, 566, "label", 360,
             color="inverse"),
        rect(556, 374, 2, 236, fill="blue"),
        text("被忽视的代价", 580, 374, "label", 620, color="gray"),
        text("01", 580, 414, "small", 40, color="blue"),
        text("时间碎片化", 626, 410, "heading", 560, color="ink"),
        text("会议打断了需要连续投入的深度工作。", 626, 448, "small", 540,
             color="gray"),
        text("02", 580, 488, "small", 40, color="blue"),
        text("问题没有被推进", 626, 484, "heading", 560, color="ink"),
        text("讨论停在现状，下一步仍然悬而未决。", 626, 522, "small", 540,
             color="gray"),
        text("03", 580, 562, "small", 40, color="blue"),
        text("结论没有责任人", 626, 558, "heading", 560, color="ink"),
        text("没有负责人，共识无法转化为行动。", 626, 596, "small", 540,
             color="gray"),
    ]
    three = frame(4, 3, "03 / PRINCIPLES", "先判断，再安排",
                  "减少无效会议 / 三条原则")
    three += [
        text("三条改进原则", 80, 126, "title", 900, color="ink", font_role="title"),
        text("把“是否必须开会”放回安排会议之前。", 80, 198, "small", 900,
             color="gray"),
        rect(120, 258, 2, 300, fill="hair"),
        text("01", 80, 262, "small", 45, color="terracotta"),
        text("能异步，就不开会", 190, 252, "title", 700, color="ink",
             font_role="title"),
        text("文档、评论与明确截止时间可以替代状态同步。", 190, 310, "small",
             700, color="gray"),
        text("02", 80, 370, "small", 45, color="terracotta"),
        text("只请必要的人", 190, 360, "title", 700, color="ink",
             font_role="title"),
        text("参会名单围绕决策者、执行者与必要信息提供者。", 190, 418, "small",
             700, color="gray"),
        text("03", 80, 478, "small", 45, color="terracotta"),
        text("没有产出，就不安排", 190, 468, "title", 700, color="ink",
             font_role="title"),
        text("会前写清议程、预期产出与结论负责人。", 190, 526, "small",
             700, color="gray"),
        oval(940, 250, 480, 480, fill="blue"),
        text("决策顺序", 1002, 330, "micro", 198, color="paper"),
        text("先问，", 1002, 354, "heading", 198, color="paper",
             font_role="title"),
        text("再开。", 1002, 406, "heading", 198, color="paper",
             font_role="title"),
        text("原则不是口号。", 1002, 486, "label", 198, color="paper"),
    ]
    four = frame(4, 4, "04 / NEXT WEEK", "2026.10.26—10.30",
                 "减少无效会议 / 下周行动")
    four += [
        text("下周行动", 80, 126, "title", 900, color="ink", font_role="title"),
        text("把三条原则放进本周日程。", 80, 198, "small", 900, color="gray"),
        text("周一", 80, 268, "action_day", 300, color="ink", font_role="title"),
        text("10.26", 80, 320, "date", 300, color="blue"),
        text("盘点本周会议", 80, 356, "body", 300, color="ink"),
        text("取消无目标项", 80, 392, "body", 300, color="gray"),
        text("周三", 420, 268, "action_day", 300, color="ink", font_role="title"),
        text("10.28", 420, 320, "date", 300, color="blue"),
        text("把一次日常同步", 420, 356, "body", 300, color="ink"),
        text("改为书面更新", 420, 392, "body", 300, color="gray"),
        text("周五", 760, 268, "action_day", 300, color="ink", font_role="title"),
        text("10.30", 760, 320, "date", 300, color="blue"),
        text("复盘会议时长", 760, 356, "body", 300, color="ink"),
        text("人数与产出", 760, 392, "body", 300, color="gray"),
        rect(80, 526, 1080, 104, fill="blue", role="outcome_band"),
        text("会议时间", 112, 542, "micro", 280, color="paper"),
        text("降低 20%", 112, 566, "outcome", 280, color="paper"),
        text("关键结论负责人覆盖", 626, 542, "micro", 500, color="paper"),
        text("100%", 626, 566, "outcome", 500, color="paper"),
    ]
    return [one, two, three, four]


def feedback_pages():
    one = [
        rect(0, 0, 36, 720, fill="terracotta", role="chrome"),
        rect(1244, 0, 36, 720, fill="blue", role="chrome"),
        oval(865, -190, 585, 585, fill="blue", role="hero_arc"),
        oval(950, -92, 415, 415, stroke="gray", sw=1),
        text("CUSTOMER EXPERIENCE / OPERATING BASELINE", 80, 62, "micro",
             1040, color="gray", semantic_role="chrome"),
        text("2026.10", 960, 62, "micro", 230, color="gray", align="right",
             semantic_role="chrome"),
        text("FEEDBACK / ROUTE / RESOLVE", 80, 154, "micro", 760,
             color="gray", semantic_role="kicker"),
        text("缩短客户", 80, 198, "hero", 760, color="ink", font_role="title",
             semantic_role="hero"),
        text("反馈闭环", 80, 314, "hero", 760, color="ink", font_role="title",
             semantic_role="hero"),
        text("6 页运营诊断 / 示例目标", 80, 492, "heading", 760, color="gray"),
        text("RECEIVE", 1012, 164, "label", 170, color="paper"),
        text("ROUTE", 1083, 312, "label", 150, color="paper"),
        text("RESOLVE", 1012, 460, "label", 170, color="paper"),
        text("客户体验 / 2026 年 10 月", 80, 680, "micro", 1000,
             color="gray", semantic_role="chrome"),
        text("01 / 06", 1040, 680, "micro", 160, color="gray",
             align="right", semantic_role="chrome"),
    ]

    two = frame(6, 2, "02 / FRICTION COST", "等待 / 转派 / 确认",
                "缩短客户反馈闭环 / 摩擦成本")
    two += [
        text("闭环慢在哪里", 80, 126, "title", 1000, color="ink", font_role="title"),
        text("反馈不是没人处理，", 80, 224, "lead", 520, color="ink", font_role="title"),
        text("而是每一步都在等。", 80, 286, "lead", 520, color="ink", font_role="title"),
        text("客户提交一次问题，团队却要", 80, 382, "body", 480, color="gray"),
        text("反复确认上下文、寻找责任人、", 80, 416, "body", 480, color="gray"),
        text("等待结论。", 80, 450, "body", 480, color="gray"),
        rect(620, 190, 580, 420, fill="panel", role="data_block"),
        text("数据对比 / 示例目标", 648, 211, "label", 330, color="gray"),
        text("现状", 1010, 211, "label", 70, color="terracotta", align="right"),
        text("目标", 1146, 211, "label", 70, color="blue", align="right"),
        text("首次回应", 648, 262, "body", 270, color="ink"),
        text("30h", 1010, 262, "body", 70, color="terracotta", align="right"),
        text("6h", 1146, 262, "body", 70, color="blue", align="right"),
        rect(648, 300, 260, 8, fill="terracotta"),
        rect(648, 318, 52, 8, fill="blue"),
        text("首次受理", 648, 368, "body", 270, color="ink"),
        text("36h", 1010, 368, "body", 70, color="terracotta", align="right"),
        text("12h", 1146, 368, "body", 70, color="blue", align="right"),
        rect(648, 406, 260, 8, fill="terracotta"),
        rect(648, 424, 87, 8, fill="blue"),
        text("闭环完成", 648, 474, "body", 270, color="ink"),
        text("5d", 1010, 474, "body", 70, color="terracotta", align="right"),
        text("2d", 1146, 474, "body", 70, color="blue", align="right"),
        rect(648, 512, 260, 8, fill="terracotta"),
        rect(648, 530, 104, 8, fill="blue"),
        text("示例目标，非实测数据。", 648, 566, "micro", 480, color="gray"),
    ]

    three = frame(6, 3, "03 / DIAGNOSTIC", "入口 / 责任 / 结论",
                  "缩短客户反馈闭环 / 断点诊断")
    three += [
        text("断点不在能力，", 80, 126, "title", 1000, color="ink", font_role="title"),
        text("而在交接。", 80, 188, "title", 1000, color="ink", font_role="title"),
        rect(80, 276, 320, 2, fill="terracotta", role="column_rule"),
        rect(462, 276, 320, 2, fill="hair", role="column_rule"),
        rect(844, 276, 320, 2, fill="hair", role="column_rule"),
        text("01", 80, 302, "number", 180, color="terracotta"),
        text("02", 462, 302, "number", 180, color="blue"),
        text("03", 844, 302, "number", 180, color="terracotta"),
        text("反馈分散", 80, 374, "heading", 320, color="ink", font_role="title"),
        text("责任漂移", 462, 374, "heading", 320, color="ink", font_role="title"),
        text("确认缺位", 844, 374, "heading", 320, color="ink", font_role="title"),
        text("表单、聊天与邮件", 80, 442, "body", 320, color="gray"),
        text("各自收集，上下文", 80, 478, "body", 320, color="gray"),
        text("无法合并。", 80, 514, "body", 320, color="gray"),
        text("问题跨团队流转，", 462, 442, "body", 320, color="gray"),
        text("却没有人持续追踪。", 462, 478, "body", 320, color="gray"),
        text("处理完成后等待", 844, 442, "body", 320, color="gray"),
        text("客户再次确认，", 844, 478, "body", 320, color="gray"),
        text("周期继续拉长。", 844, 514, "body", 320, color="gray"),
        rect(80, 604, 1080, 44, fill="terracotta"),
        text("先统一入口，再定义责任，最后锁定关闭条件。", 104, 612, "body",
             1030, color="paper"),
    ]

    four = frame(6, 4, "04 / THREE DAY LOOP", "接收 / 处理 / 关闭",
                 "缩短客户反馈闭环 / 目标流程")
    four += [
        text("把闭环压进三天", 80, 126, "title", 1000, color="ink", font_role="title"),
        text("用时间边界替换口头承诺。", 80, 198, "heading", 900, color="gray"),
        rect(80, 270, 336, 270, fill="panel"),
        rect(472, 270, 336, 270, fill="panel_blue"),
        rect(864, 270, 296, 270, fill="panel"),
        rect(110, 410, 1020, 3, fill="terracotta"),
        oval(154, 394, 32, 32, fill="terracotta"),
        oval(586, 394, 32, 32, fill="terracotta"),
        oval(950, 394, 32, 32, fill="terracotta"),
        text("T+0", 110, 292, "heading", 200, color="terracotta"),
        text("统一接收", 110, 342, "heading", 250, color="ink", font_role="title"),
        text("客户服务台", 110, 448, "body", 260, color="ink"),
        text("合并入口，补齐环境、", 110, 486, "small", 300, color="gray"),
        text("影响与期望。", 110, 516, "small", 300, color="gray"),
        text("T+1", 502, 292, "heading", 200, color="blue"),
        text("明确责任", 502, 342, "heading", 250, color="ink", font_role="title"),
        text("产品责任人", 502, 448, "body", 260, color="ink"),
        text("完成分类，给出处理路径", 502, 486, "small", 300, color="gray"),
        text("与更新时间。", 502, 516, "small", 300, color="gray"),
        text("T+2", 894, 292, "heading", 190, color="terracotta"),
        text("验证关闭", 894, 342, "heading", 240, color="ink", font_role="title"),
        text("交付负责人", 894, 448, "body", 240, color="ink"),
        text("对照验收样例，记录结论", 894, 486, "small", 300, color="gray"),
        text("与剩余风险。", 894, 516, "small", 300, color="gray"),
        text("时间边界", 80, 572, "label", 180, color="gray"),
        text("48 小时形成可验证的关闭结论", 248, 566, "heading", 900,
             color="ink", font_role="title"),
    ]

    five = frame(6, 5, "05 / OWNERSHIP", "角色 / 输入 / 输出",
                 "缩短客户反馈闭环 / 责任界面")
    five += [
        oval(-260, 370, 650, 650, fill="blue", role="background_arc"),
        text("每个交接点", 80, 126, "title", 1000, color="ink", font_role="title"),
        text("都要有名字。", 80, 188, "title", 1000, color="ink", font_role="title"),
        text("责任人不是审批人，", 80, 294, "heading", 320, color="paper"),
        text("而是确保下一步", 80, 336, "heading", 320, color="paper"),
        text("按时发生的人。", 80, 378, "heading", 320, color="paper"),
        rect(620, 180, 580, 440, fill="panel"),
        text("角色", 648, 204, "label", 140, color="gray"),
        text("输入端", 790, 204, "label", 180, color="gray"),
        text("输出端", 1004, 204, "label", 160, color="gray"),
        text("客服台", 648, 264, "body", 130, color="ink"),
        text("多渠道反馈", 790, 264, "body", 190, color="ink"),
        text("统一工单", 1004, 264, "body", 160, color="ink"),
        text("产品负责人", 648, 356, "body", 170, color="ink"),
        text("工单与影响", 790, 356, "body", 190, color="ink"),
        text("处理路径", 1004, 356, "body", 160, color="ink"),
        text("交付负责人", 648, 448, "body", 170, color="ink"),
        text("验收样例", 790, 448, "body", 190, color="ink"),
        text("关闭结论", 1004, 448, "body", 160, color="ink"),
        text("一条反馈只有一个主责任人。", 648, 554, "body", 520,
             color="terracotta"),
    ]

    six = frame(6, 6, "06 / THIS WEEK", "2026.10.26—10.30",
                "缩短客户反馈闭环 / 本周启动")
    six += [
        text("本周启动", 80, 126, "title", 1000, color="ink", font_role="title"),
        text("先跑通一条真实反馈，再扩到全部入口。", 80, 198, "heading",
             1000, color="gray"),
        text("周一", 80, 262, "action_day", 300, color="ink", font_role="title"),
        text("10.26", 80, 320, "date", 200, color="blue"),
        text("选一条跨团队反馈", 80, 360, "body", 300, color="ink"),
        text("记录每次等待与转派", 80, 406, "small", 300, color="gray"),
        text("周三", 480, 262, "action_day", 280, color="ink", font_role="title"),
        text("10.28", 480, 320, "date", 200, color="blue"),
        text("指定唯一主责任人", 480, 360, "body", 280, color="ink"),
        text("公开处理路径与更新时间", 480, 406, "small", 300, color="gray"),
        text("周五", 840, 262, "action_day", 260, color="ink", font_role="title"),
        text("10.30", 840, 320, "date", 200, color="blue"),
        text("对照验收样例关闭", 840, 360, "body", 300, color="ink"),
        text("复盘等待点与遗漏信息", 840, 406, "small", 300, color="gray"),
        rect(80, 560, 1080, 78, fill="blue", role="outcome_band"),
        text("首次回应", 112, 572, "label", 180, color="paper"),
        text("≤ 6h", 112, 598, "heading", 220, color="paper"),
        text("闭环完成", 494, 572, "label", 180, color="paper"),
        text("≤ 2d", 494, 598, "heading", 220, color="paper"),
        text("主责任人覆盖率", 846, 572, "label", 240, color="paper"),
        text("100%", 846, 598, "heading", 220, color="paper"),
        text("示例目标，非实测数据。", 80, 646, "micro", 900, color="gray"),
    ]
    return [one, two, three, four, five, six]


def profile_style(source_pages, style, profile_name):
    profile = style["profiles"][profile_name]
    palette = profile["palette"]
    text_map = profile.get("text_color_map", {})
    fill_map = profile.get("fill_color_map", {})
    stroke_map = profile.get("stroke_color_map", {})
    omitted = set(profile.get("omit_shape_roles", []))
    default_background = palette.get("background", palette.get("paper"))
    resolved_pages = []
    for source_page in source_pages:
        resolved = [{
            "kind": "rect",
            "x": 0,
            "y": 0,
            "w": style["grid"]["canvas_px"][0],
            "h": style["grid"]["canvas_px"][1],
            "fill": default_background,
            "stroke": None,
            "stroke_width_px": 0,
            "role": "background",
            "name": "Style background",
        }]
        for source in source_page:
            if source["kind"] == "text":
                item = deepcopy(source)
                item["size"] = profile["type_by_role"][item["type_role"]]
                color_role = text_map.get(item["color"], item["color"])
                item["color_hex"] = palette[color_role]
                item["font"] = style["fonts"][item["font_role"]]
                resolved.append(item)
            else:
                if source["role"] in omitted:
                    continue
                item = deepcopy(source)
                fill_role = fill_map.get(item["fill"], item["fill"]) if item["fill"] else None
                stroke_role = stroke_map.get(item["stroke"], item["stroke"]) if item["stroke"] else None
                item["fill"] = palette[fill_role] if fill_role else None
                item["stroke"] = palette[stroke_role] if stroke_role else None
                resolved.append(item)
        resolved_pages.append(resolved)
    return resolved_pages


def add_text(slide, item, style):
    shape = slide.shapes.add_textbox(
        Emu(round(item["x"] * 9525)),
        Emu(round(item["y"] * 9525)),
        Emu(round(item["w"] * 9525)),
        Emu(round(item["h"] * 9525)),
    )
    shape.name = item["name"]
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = False
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = {
        "left": PP_ALIGN.LEFT,
        "right": PP_ALIGN.RIGHT,
        "center": PP_ALIGN.CENTER,
    }[item["align"]]
    paragraph.space_before = paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = item["text"]
    run.font.name = item["font"]
    run.font.size = Pt(item["size"] * style["fonts"]["pixels_to_points"])
    run.font.color.rgb = RGBColor.from_string(item["color_hex"])
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", item["font"])


def add_shape(slide, item, style):
    shape_type = MSO_SHAPE.OVAL if item["kind"] == "oval" else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type,
        Emu(round(item["x"] * 9525)),
        Emu(round(item["y"] * 9525)),
        Emu(round(item["w"] * 9525)),
        Emu(round(item["h"] * 9525)),
    )
    shape.name = item["name"]
    shape.shadow.inherit = False
    element_style = shape._element.find(qn("p:style"))
    if element_style is not None:
        shape._element.remove(element_style)
    if item["fill"]:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(item["fill"])
    else:
        shape.fill.background()
    if item["stroke"]:
        shape.line.color.rgb = RGBColor.from_string(item["stroke"])
        shape.line.width = Pt(
            item["stroke_width_px"] * style["shapes"]["stroke_points_factor"]
        )
    else:
        shape.line.fill.background()


def clear_slides(prs):
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.get(qn("r:id")))
        slide_ids.remove(slide_id)


def build_presentation(pages, style):
    prs = Presentation(DONOR)
    clear_slides(prs)
    prs.slide_width = Emu(round(style["grid"]["canvas_px"][0] * 9525))
    prs.slide_height = Emu(round(style["grid"]["canvas_px"][1] * 9525))
    for page in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for item in page:
            if item["kind"] == "text":
                add_text(slide, item, style)
            else:
                add_shape(slide, item, style)
    return prs


def subset_fonts(src, dst, allowed_families):
    with zipfile.ZipFile(src) as source:
        presentation = etree.fromstring(source.read("ppt/presentation.xml"))
        rels = etree.fromstring(source.read("ppt/_rels/presentation.xml.rels"))
        font_lst = presentation.find(".//p:embeddedFontLst", NS)
        kept_ids = set()
        if font_lst is not None:
            for embedded_font in list(font_lst):
                typeface = embedded_font.find("p:font", NS).get("typeface")
                if typeface not in allowed_families:
                    font_lst.remove(embedded_font)
                    continue
                for child in embedded_font:
                    rid = child.get(qn("r:id"))
                    if rid:
                        kept_ids.add(rid)
        targets = {
            relationship.get("Id"): relationship.get("Target")
            for relationship in rels.findall("r:Relationship", REL_NS)
        }
        removed_parts = set()
        for relationship in list(rels):
            if relationship.get("Type", "").endswith("/font"):
                if relationship.get("Id") not in kept_ids:
                    target = relationship.get("Target", "")
                    if target:
                        removed_parts.add("ppt/" + target.lstrip("./"))
                    rels.remove(relationship)
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                if info.filename in removed_parts:
                    continue
                data = source.read(info.filename)
                if info.filename == "ppt/presentation.xml":
                    data = etree.tostring(
                        presentation, encoding="UTF-8",
                        xml_declaration=True, standalone=True
                    )
                elif info.filename == "ppt/_rels/presentation.xml.rels":
                    data = etree.tostring(
                        rels, encoding="UTF-8",
                        xml_declaration=True, standalone=True
                    )
                target.writestr(info, data)


def make_timing(targets, rule):
    timing = etree.Element(qn("p:timing"))
    tn_lst = etree.SubElement(timing, qn("p:tnLst"))
    root_par = etree.SubElement(tn_lst, qn("p:par"))
    root_ctn = etree.SubElement(root_par, qn("p:cTn"))
    root_ctn.set("id", "1")
    root_ctn.set("dur", "indefinite")
    root_ctn.set("restart", "never")
    root_ctn.set("nodeType", "tmRoot")
    seq = etree.SubElement(etree.SubElement(root_ctn, qn("p:childTnLst")), qn("p:seq"))
    seq.set("concurrent", "1")
    seq.set("nextAc", "seek")
    main = etree.SubElement(seq, qn("p:cTn"))
    main.set("id", "2")
    main.set("dur", "indefinite")
    main.set("nodeType", "mainSeq")
    group = etree.SubElement(etree.SubElement(main, qn("p:childTnLst")), qn("p:par"))
    group_ctn = etree.SubElement(group, qn("p:cTn"))
    group_ctn.set("id", "3")
    group_ctn.set("fill", "hold")
    group_start = etree.SubElement(group_ctn, qn("p:stCondLst"))
    cond = etree.SubElement(group_start, qn("p:cond"))
    cond.set("delay", "indefinite")
    begin = etree.SubElement(group_start, qn("p:cond"))
    begin.set("evt", "onBegin")
    begin.set("delay", "0")
    etree.SubElement(begin, qn("p:tn")).set("val", "2")
    overlap = etree.SubElement(etree.SubElement(group_ctn, qn("p:childTnLst")), qn("p:par"))
    overlap_ctn = etree.SubElement(overlap, qn("p:cTn"))
    overlap_ctn.set("id", "4")
    overlap_ctn.set("fill", "hold")
    etree.SubElement(etree.SubElement(overlap_ctn, qn("p:stCondLst")), qn("p:cond")).set("delay", "0")
    children = etree.SubElement(overlap_ctn, qn("p:childTnLst"))

    next_id = 5
    for index, shape_id in enumerate(targets):
        par = etree.SubElement(children, qn("p:par"))
        ctn = etree.SubElement(par, qn("p:cTn"))
        ctn.set("id", str(next_id))
        ctn.set("presetID", "10")
        ctn.set("presetClass", "entr")
        ctn.set("presetSubtype", "0")
        ctn.set("fill", "hold")
        ctn.set("nodeType", "withEffect")
        start = etree.SubElement(ctn, qn("p:stCondLst"))
        etree.SubElement(start, qn("p:cond")).set(
            "delay", str(index * rule["stagger_ms"])
        )
        effect_children = etree.SubElement(ctn, qn("p:childTnLst"))
        set_effect = etree.SubElement(effect_children, qn("p:set"))
        behavior = etree.SubElement(set_effect, qn("p:cBhvr"))
        behavior_ctn = etree.SubElement(behavior, qn("p:cTn"))
        behavior_ctn.set("id", str(next_id + 1))
        behavior_ctn.set("dur", "1")
        etree.SubElement(etree.SubElement(behavior_ctn, qn("p:stCondLst")), qn("p:cond")).set("delay", "0")
        target_node = etree.SubElement(behavior, qn("p:tgtEl"))
        etree.SubElement(target_node, qn("p:spTgt")).set("spid", str(shape_id))
        attrs = etree.SubElement(behavior, qn("p:attrNameLst"))
        etree.SubElement(attrs, qn("p:attrName")).text = "style.visibility"
        etree.SubElement(etree.SubElement(set_effect, qn("p:to")), qn("p:strVal")).set("val", "visible")
        fade = etree.SubElement(effect_children, qn("p:animEffect"))
        fade.set("transition", rule["direction"])
        fade.set("filter", rule["effect"])
        fade_behavior = etree.SubElement(fade, qn("p:cBhvr"))
        fade_ctn = etree.SubElement(fade_behavior, qn("p:cTn"))
        fade_ctn.set("id", str(next_id + 2))
        fade_ctn.set("dur", str(rule["duration_ms"]))
        fade_target = etree.SubElement(fade_behavior, qn("p:tgtEl"))
        etree.SubElement(fade_target, qn("p:spTgt")).set("spid", str(shape_id))
        next_id += 3
    previous = etree.SubElement(etree.SubElement(seq, qn("p:prevCondLst")), qn("p:cond"))
    previous.set("evt", "onPrev")
    previous.set("delay", "0")
    etree.SubElement(etree.SubElement(previous, qn("p:tgtEl")), qn("p:sldTgt"))
    following = etree.SubElement(etree.SubElement(seq, qn("p:nextCondLst")), qn("p:cond"))
    following.set("evt", "onNext")
    following.set("delay", "0")
    etree.SubElement(etree.SubElement(following, qn("p:tgtEl")), qn("p:sldTgt"))
    build_list = etree.SubElement(timing, qn("p:bldLst"))
    for shape_id in targets:
        build_shape = etree.SubElement(build_list, qn("p:bldP"))
        build_shape.set("spid", str(shape_id))
        build_shape.set("grpId", "0")
    return timing


def add_animations(src, dst, rule, counts):
    report = {}
    with zipfile.ZipFile(src) as source:
        presentation = etree.fromstring(source.read("ppt/presentation.xml"))
        rels = etree.fromstring(source.read("ppt/_rels/presentation.xml.rels"))
        targets_map = {
            rel.get("Id"): rel.get("Target")
            for rel in rels.findall("r:Relationship", REL_NS)
        }
        replacements = {}
        slide_ids = presentation.xpath("//p:sldIdLst/p:sldId", namespaces=NS)
        for page, slide_id in enumerate(slide_ids, 1):
            target = targets_map[slide_id.get(qn("r:id"))]
            part = target.lstrip("/") if target.startswith("/") else "ppt/" + target
            root = etree.fromstring(source.read(part))
            text_ids = [
                int(value) for value in root.xpath(
                    ".//p:sp[.//a:t]/p:nvSpPr/p:cNvPr/@id", namespaces=NS
                )
            ]
            all_ids = [
                int(value) for value in root.xpath(
                    "./p:cSld/p:spTree/p:sp/p:nvSpPr/p:cNvPr/@id", namespaces=NS
                )
            ]
            ordered = text_ids + [item for item in all_ids if item not in text_ids]
            targets = ordered[:counts[page - 1]]
            old = root.find("p:timing", NS)
            if old is not None:
                old.getparent().remove(old)
            timing = make_timing(targets, rule)
            ext = root.find("p:extLst", NS)
            if ext is not None:
                ext.addprevious(timing)
            else:
                root.append(timing)
            replacements[part] = etree.tostring(
                root, encoding="UTF-8", xml_declaration=True, standalone=True
            )
            report[page] = len(root.xpath(".//p:animEffect", namespaces=NS))
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as target_archive:
            for info in source.infolist():
                target_archive.writestr(info, replacements.get(info.filename, source.read(info.filename)))
    return report


def render_svg(page, profile, style, target):
    canvas_w, canvas_h = style["grid"]["canvas_px"]
    background = profile["palette"].get("background", profile["palette"].get("paper"))
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{canvas_w}" height="{canvas_h}" viewBox="0 0 {canvas_w} {canvas_h}">',
        f'<rect width="{canvas_w}" height="{canvas_h}" fill="#{background}"/>',
    ]
    for item in page:
        if item["kind"] == "rect":
            fill = f'fill="#{item["fill"]}"' if item["fill"] else 'fill="none"'
            stroke = f'stroke="#{item["stroke"]}"' if item["stroke"] else 'stroke="none"'
            parts.append(
                f'<rect x="{item["x"]}" y="{item["y"]}" width="{item["w"]}" height="{item["h"]}" '
                f'{fill} {stroke} stroke-width="{item["stroke_width_px"]}"/>'
            )
        elif item["kind"] == "oval":
            fill = f'fill="#{item["fill"]}"' if item["fill"] else 'fill="none"'
            stroke = f'stroke="#{item["stroke"]}"' if item["stroke"] else 'stroke="none"'
            parts.append(
                f'<ellipse cx="{item["x"] + item["w"] / 2}" cy="{item["y"] + item["h"] / 2}" '
                f'rx="{item["w"] / 2}" ry="{item["h"] / 2}" '
                f'{fill} {stroke} stroke-width="{item["stroke_width_px"]}"/>'
            )
        else:
            x = item["x"] + (item["w"] if item["align"] == "right" else item["w"] / 2 if item["align"] == "center" else 0)
            anchor = "end" if item["align"] == "right" else "middle" if item["align"] == "center" else "start"
            family = style["fonts"]["svg_families"][item["font_role"]]
            parts.append(
                f'<text x="{x}" y="{item["y"] + item["size"] * 0.84}" fill="#{item["color_hex"]}" '
                f'font-family="{family}" font-size="{item["size"]}" text-anchor="{anchor}">'
                f'{html.escape(item["text"])}</text>'
            )
    parts.append("</svg>")
    target.write_text("\n".join(parts), encoding="utf-8")


def render_pngs(pages, profile, style, output, temp):
    svg_dir = temp / "svg"
    html_dir = temp / "html"
    preview = output / "preview"
    for directory in (svg_dir, html_dir, preview):
        directory.mkdir(parents=True, exist_ok=True)
    background = profile["palette"].get("background", profile["palette"].get("paper"))
    for page_number, page in enumerate(pages, 1):
        svg = svg_dir / f"slide_{page_number:02d}.svg"
        render_svg(page, profile, style, svg)
        page_html = html_dir / f"s{page_number:02d}.html"
        page_html.write_text(
            "<!doctype html><html><head><meta charset='utf-8'>"
            "<style>html,body{margin:0;width:1280px;height:720px;overflow:hidden;"
            f"background:#{background}}}img{{display:block;width:1280px;height:720px}}</style>"
            f"</head><body><img src='../_svg/{svg.name}'></body></html>",
            encoding="utf-8",
        )
        png = preview / f"s{page_number:02d}.png"
        subprocess.run(
            [
                str(EDGE), "--headless=old", "--disable-gpu",
                "--disable-gpu-compositing", "--in-process-gpu", "--no-sandbox",
                "--hide-scrollbars", "--force-device-scale-factor=1",
                "--window-size=1280,720",
                f"--user-data-dir={temp / ('edge_' + str(page_number))}",
                f"--screenshot={png}", page_html.as_uri(),
            ],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
    return [preview / f"s{page:02d}.png" for page in range(1, len(pages) + 1)]


def make_contact(pngs, target, labels):
    thumbs = []
    for path in pngs:
        with Image.open(path) as image:
            thumbs.append(image.convert("RGB").resize((960, 540), Image.Resampling.LANCZOS))
    rows = (len(thumbs) + 1) // 2
    width = 2008
    height = 20 + rows * 608 + 16
    sheet = Image.new("RGB", (width, height), (244, 241, 233))
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.truetype(r"C:\Windows\Fonts\simkai.ttf", 30)
    for index, (thumb, label) in enumerate(zip(thumbs, labels)):
        row, col = divmod(index, 2)
        x, y = 20 + col * 1008, 20 + row * 608
        draw.text((x, y), label, fill=(30, 30, 30), font=font)
        sheet.paste(thumb, (x, y + 44))
    sheet.save(target, optimize=True)


def check_text(pages):
    issues = []
    for page_number, page in enumerate(pages, 1):
        for item in page:
            if item["kind"] != "text":
                continue
            font_path = Path(
                r"C:\Windows\Fonts\simkai.ttf"
                if item["font"] == "楷体" else
                r"C:\Windows\Fonts\simsun.ttc"
            )
            font = ImageFont.truetype(str(font_path), round(item["size"]))
            if font.getlength(item["text"]) > item["w"] + 2:
                issues.append((page_number, item["text"]))
    return issues


def build_case(name, source_pages, profile_name, expected_effects, labels):
    style = json.loads(STYLE_PATH.read_text(encoding="utf-8"))
    profile = style["profiles"][profile_name]
    pages = profile_style(source_pages, style, profile_name)
    output = OUTPUT / name
    output.mkdir(parents=True, exist_ok=True)
    final_name = "减少无效会议_4页_重建版.pptx" if name == "meeting_4page" else (
        "缩短客户反馈闭环_6页_重建版.pptx" if name == "feedback_6page"
        else "缩短客户反馈闭环_4页_深色重建版.pptx"
    )
    final_pptx = output / final_name
    with tempfile.TemporaryDirectory(dir=BASE) as temp_name:
        temp = Path(temp_name)
        raw = temp / "raw.pptx"
        embedded = temp / "embedded.pptx"
        build_presentation(pages, style).save(raw)
        subset_fonts(raw, embedded, {style["fonts"]["title"], style["fonts"]["body"]})
        counts = add_animations(embedded, final_pptx, style["animation"], expected_effects)
        pngs = render_pngs(pages, profile, style, output, temp)
    contact = output / "contact_sheet.png"
    make_contact(pngs, contact, labels)
    issues = check_text(pages)
    markers = 0
    round_corners = 0
    gradients = 0
    page_accents = {}
    allowed_accents = {
        profile["palette"].get(name, "")
        for name in (("terracotta", "blue") if profile_name == "swiss_default" else ("accent",))
        if name in profile["palette"] and profile["palette"].get(name)
    }
    with zipfile.ZipFile(final_pptx) as archive:
        presentation = etree.fromstring(archive.read("ppt/presentation.xml"))
        rels = etree.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
        targets = {
            rel.get("Id"): rel.get("Target")
            for rel in rels.findall("r:Relationship", REL_NS)
        }
        markers = archive.read("ppt/presentation.xml").decode("utf-8").count("embeddedFont")
        for page_number, slide_id in enumerate(
            presentation.xpath("//p:sldIdLst/p:sldId", namespaces=NS), 1
        ):
            target = targets[slide_id.get(qn("r:id"))]
            part = target.lstrip("/") if target.startswith("/") else "ppt/" + target
            root = etree.fromstring(archive.read(part))
            round_corners += len(root.xpath(".//a:prstGeom[@prst='roundRect']", namespaces=NS))
            gradients += len(root.xpath(".//a:gradFill", namespaces=NS))
            page_accents[page_number] = sorted({
                value.upper()
                for value in root.xpath(".//a:srgbClr/@val", namespaces=NS)
                if value.upper() in allowed_accents
            })
    black_images = []
    for path in pngs:
        with Image.open(path) as image:
            colors = image.convert("RGB").getcolors(2)
        if colors is not None and len(colors) == 1 and colors[0][1] == (0, 0, 0):
            black_images.append(str(path))
    effect_values = list(counts.values())
    acceptance = {
        "status": "PASS" if (
            not issues
            and effect_values == expected_effects
            and markers == 6
            and round_corners == 0
            and gradients == 0
            and all(len(values) <= 2 for values in page_accents.values())
            and not black_images
        ) else "FAIL",
        "name": name,
        "pages": len(pages),
        "text_overflows": issues,
        "font_embedded_markers": markers,
        "effect_counts": effect_values,
        "expected_effect_counts": expected_effects,
        "round_corner_count": round_corners,
        "gradient_count": gradients,
        "accent_colors_per_page": page_accents,
        "black_images": black_images,
        "pptx": str(final_pptx),
        "contact_sheet": str(contact),
        "pngs": [str(path) for path in pngs],
        "sha256": hashlib.sha256(final_pptx.read_bytes()).hexdigest(),
    }
    (output / "acceptance.json").write_text(
        json.dumps(acceptance, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return acceptance


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    meeting = meeting_pages()
    feedback = feedback_pages()
    results = [
        build_case("meeting_4page", meeting, "swiss_default", [9, 21, 22, 24],
                   ["S01 封面", "S02 会议成本", "S03 改进原则", "S04 下周行动"]),
        build_case("feedback_6page", feedback, "swiss_default", [11, 23, 23, 23, 22, 25],
                   ["S01 封面", "S02 摩擦成本", "S03 断点诊断", "S04 三天流程", "S05 责任界面", "S06 本周启动"]),
        build_case("feedback_4page_dark", feedback[:4], "dark_editorial", [11, 23, 23, 23],
                   ["S01 封面", "S02 摩擦成本", "S03 断点诊断", "S04 三天流程"]),
    ]
    status = "PASS" if all(item["status"] == "PASS" for item in results) else "FAIL"
    (BASE / "rebuild_verification.json").write_text(
        json.dumps({"status": status, "cases": results}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps({"status": status, "cases": results}, ensure_ascii=False, indent=2))
    if status != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
