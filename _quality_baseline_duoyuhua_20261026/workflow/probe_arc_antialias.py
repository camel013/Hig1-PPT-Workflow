from __future__ import annotations

import json
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu


BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "_arc_antialias_probe"
PAPER = (244, 241, 233)
BLUE = (90, 112, 126)


def draw_circle(size: tuple[int, int], box: tuple[int, int, int, int]) -> Image.Image:
    image = Image.new("RGB", size, PAPER)
    ImageDraw.Draw(image).ellipse(box, fill=BLUE)
    return image


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    direct = draw_circle((1280, 720), (400, 120, 880, 600))
    direct.save(OUT / "direct_1x.png")

    scale = 4
    supersampled = draw_circle(
        (1280 * scale, 720 * scale),
        tuple(value * scale for value in (400, 120, 880, 600)),
    )
    antialiased = supersampled.resize((1280, 720), Image.Resampling.LANCZOS)
    antialiased.save(OUT / "supersampled_4x.png")

    presentation = Presentation()
    presentation.slide_width = Emu(round(1280 * 9525))
    presentation.slide_height = Emu(round(720 * 9525))
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(*PAPER)
    background.line.fill.background()
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Emu(round(400 * 9525)),
        Emu(round(120 * 9525)),
        Emu(round(480 * 9525)),
        Emu(round(480 * 9525)),
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = RGBColor(*BLUE)
    circle.line.fill.background()
    presentation.save(OUT / "native_circle.pptx")
    print(json.dumps({
        "direct": str(OUT / "direct_1x.png"),
        "supersampled": str(OUT / "supersampled_4x.png"),
        "pptx": str(OUT / "native_circle.pptx"),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
