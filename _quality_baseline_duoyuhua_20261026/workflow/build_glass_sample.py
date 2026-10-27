from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu


BASE = Path(__file__).resolve().parent.parent
WORKFLOW = Path(__file__).resolve().parent
sys.path.insert(0, str(WORKFLOW))
import rebuild_all as pipeline


STYLE = json.loads(pipeline.STYLE_PATH.read_text(encoding="utf-8"))
PROFILE = STYLE["profiles"]["swiss_default"]
PALETTE = PROFILE["palette"]
OUTPUT = BASE.parent / "glass_layer_sample_20261026"
SLIDE = OUTPUT / "玻璃质感色块_隔离样页.pptx"
BACKGROUND = OUTPUT / "_glass_background.png"


def color(name: str) -> tuple[int, int, int]:
    value = PALETTE[name]
    return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))


def px(value: float) -> Emu:
    return Emu(round(value * 9525))


def set_alpha(shape, alpha_pct: int):
    solid = shape._element.spPr.find(pipeline.qn("a:solidFill"))
    srgb = solid.find(pipeline.qn("a:srgbClr"))
    alpha = pipeline.etree.SubElement(srgb, pipeline.qn("a:alpha"))
    alpha.set("val", str(alpha_pct * 1000))


def make_background():
    scale = 4
    image = Image.new("RGB", (1280 * scale, 720 * scale), color("paper"))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 620 * scale, 720 * scale), fill=color("terracotta"))
    draw.ellipse(
        (760 * scale, -120 * scale, 1320 * scale, 440 * scale),
        fill=color("blue"),
    )
    draw.ellipse(
        (852 * scale, -28 * scale, 1228 * scale, 348 * scale),
        outline=color("panel_blue"),
        width=2 * scale,
    )
    for index in range(7):
        y = 350 + index * 32
        draw.rectangle(
            (540 * scale, y * scale, 1030 * scale, (y + 3) * scale),
            fill=color("gray"),
        )
    draw.rectangle(
        (88 * scale, 540 * scale, 520 * scale, 548 * scale),
        fill=color("paper"),
    )
    draw.rectangle(
        (1000 * scale, 520 * scale, 1008 * scale, 660 * scale),
        fill=color("terracotta"),
    )
    image = image.resize((1280, 720), Image.Resampling.LANCZOS)

    glass_box = (360, 180, 1080, 570)
    blurred = image.crop(glass_box).filter(ImageFilter.GaussianBlur(18))
    image.paste(blurred, glass_box)
    image.save(BACKGROUND, optimize=True)


def build():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    make_background()
    with tempfile.TemporaryDirectory(dir=BASE) as temp_dir:
        temp = Path(temp_dir)
        raw = temp / "raw.pptx"
        embedded = temp / "embedded.pptx"

        presentation = Presentation(pipeline.DONOR)
        pipeline.clear_slides(presentation)
        presentation.slide_width = px(1280)
        presentation.slide_height = px(720)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(
            str(BACKGROUND), px(0), px(0), px(1280), px(720)
        )

        glass = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, px(360), px(180), px(720), px(390)
        )
        glass.name = "glass_field"
        glass.shadow.inherit = False
        style = glass._element.find(pipeline.qn("p:style"))
        if style is not None:
            glass._element.remove(style)
        glass.fill.solid()
        glass.fill.fore_color.rgb = RGBColor.from_string(PALETTE["blue"])
        set_alpha(glass, 30)
        glass.line.fill.background()

        top = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, px(360), px(180), px(720), px(2)
        )
        top.fill.solid()
        top.fill.fore_color.rgb = RGBColor(255, 255, 255)
        set_alpha(top, 32)
        top.line.fill.background()
        left = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, px(360), px(180), px(2), px(390)
        )
        left.fill.solid()
        left.fill.fore_color.rgb = RGBColor(255, 255, 255)
        set_alpha(left, 24)
        left.line.fill.background()

        items = [
            pipeline.text("GLASS LAYER / SAMPLE", 400, 214, "micro", 500,
                          color="inverse", semantic_role="kicker"),
            pipeline.text("让色块拥有", 400, 256, "title", 430,
                          color="inverse", font_role="title"),
            pipeline.text("第二层", 400, 314, "title", 430,
                          color="inverse", font_role="title"),
            pipeline.text("背景采样，局部模糊", 400, 392, "body", 470,
                          color="inverse"),
            pipeline.text("薄层着色，单侧高光", 400, 426, "body", 470,
                          color="inverse"),
            pipeline.text("不加阴影，不做圆角", 400, 482, "small", 470,
                          color="inverse"),
        ]
        for item in items:
            item["size"] = PROFILE["type_by_role"][item["type_role"]]
            item["color_hex"] = PALETTE[item["color"]]
            item["font"] = STYLE["fonts"][item["font_role"]]
            pipeline.add_text(slide, item, STYLE)

        layout_issues = []
        boxes = []
        for item in items:
            font_path = Path(
                r"C:\Windows\Fonts\simkai.ttf"
                if item["font_role"] == "title"
                else r"C:\Windows\Fonts\simsun.ttc"
            )
            font = ImageFont.truetype(str(font_path), round(item["size"]))
            advance = font.getlength(item["text"])
            if advance > item["w"] + 2:
                layout_issues.append(
                    {"type": "overflow", "text": item["text"], "width": advance}
                )
            bbox = font.getbbox(item["text"])
            boxes.append({
                "text": item["text"],
                "x0": item["x"] + bbox[0],
                "y0": item["y"] + bbox[1],
                "x1": item["x"] + bbox[2],
                "y1": item["y"] + bbox[3],
            })
        for index, first in enumerate(boxes):
            for second in boxes[index + 1:]:
                if (
                    first["x1"] > second["x0"] + 2
                    and second["x1"] > first["x0"] + 2
                    and first["y1"] > second["y0"] + 2
                    and second["y1"] > first["y0"] + 2
                ):
                    layout_issues.append({
                        "type": "overlap",
                        "first": first["text"],
                        "second": second["text"],
                    })

        presentation.save(raw)
        pipeline.subset_fonts(
            raw, embedded, {STYLE["fonts"]["title"], STYLE["fonts"]["body"]}
        )
        presentation = Presentation(embedded)
        presentation.save(SLIDE)
    import embed_full_fonts_from_donor
    embed_full_fonts_from_donor.embed(
        SLIDE, pipeline.DONOR, {"楷体", "华文中宋"}
    )

    result = {
        "status": "PASS" if not layout_issues else "FAIL",
        "pptx": str(SLIDE),
        "layout_issues": layout_issues,
        "text_items": len(items),
        "glass": {
            "blur_radius_px": 18,
            "background_antialias_scale": 4,
            "tint": "5A707E",
            "opacity_pct": 30,
            "top_highlight_pct": 32,
            "left_highlight_pct": 24,
        },
    }
    (OUTPUT / "layout_acceptance.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if layout_issues:
        raise SystemExit(1)


if __name__ == "__main__":
    build()
